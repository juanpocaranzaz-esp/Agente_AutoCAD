"""
AutoCAD 2021 Engineering Agent MCP Server (v1.4.9)
Added in v1.4.9 (draw_primitives_batch, symmetric topology floor, T-junction detection, WCS normal vectors, scale_z mirror preservation, window_bbox fallback propagation, idempotent UCS rollback) (100% backward-compatible with v1.4.7 4-tool architecture):
- Spatial Bounding Box (bbox) & Field Projection (fields) in inspect_cad_headless.
- Pure-AutoLISP Metric Takeoffs (include_takeoff, export_takeoff_path to CSV/XLSX).
- Revision Diffing (compare_with_dwg supporting 'latest_backup' and 'original_backup').
- Automated CAD Quality & Hygiene Critique (include_critique with 0-100 score).
- Headless Visual Verification via non-modal _.-PLOT to PDF/PNG (export_plot_path).
Architecture:
- Primary Mode: Headless DWG processing via accoreconsole.exe (/readonly, 1-4s latency). Standard AutoLISP only.
- Secondary Mode: COM Automation via win32com (AutoCAD.Application.24) when user explicitly requests live GUI interaction.
- Sandboxing & Delimiters: Encloses extracted drawing texts and block attribute values in <untrusted_cad_data> XML tags while keeping CAD symbol table identifiers clean.
- Circuit Breaker: Sanitizes AutoLISP against OS calls while allowing verified .agents/skills/ and lisp-library/ scripts.
- Persistence Integrity: Pre-mutation validation, CMDACTIVE check, DBMOD persistence verification, and atomic rollback.
- Resilient Encoding: Unicode UTF-16LE / CP1252 output decoding with odd-byte misalignment handling.
- Sequential Backups: <Name>-Legacy-0.dwg, -Legacy-1.dwg... with automatic cleanup on failure.
- File Status & Locks: Verifies real Windows lock state and owner via .dwl/.dwl2 files.
- Git Synchronization: Safe startup fast-forward pull and pre-save rebase-before-push sync with human confirmation.
"""

import os
import sys
import time
import uuid
import stat
import json
import shutil
import logging
import threading
import subprocess
import csv
import html
import math
import re
from functools import wraps
from pathlib import Path
from typing import Dict, Any, Optional, List, Tuple
import xml.etree.ElementTree as ET

# FastMCP / win32 imports
from mcp.server.fastmcp import FastMCP
import win32com.client
import pythoncom
import pywintypes
import win32file

# AutoCAD COM ProgID and executable paths
ACAD_PROG_ID = "AutoCAD.Application.24"
ACAD_GENERIC_PROG_ID = "AutoCAD.Application"


def get_accoreconsole_path() -> str:
    """
    Dynamically discovers the accoreconsole.exe executable path:
    1. Checks ACAD_CONSOLE_PATH environment variable if set and exists.
    2. Probes standard Autodesk installation directories across C: and D: drives (AutoCAD 2021 through 2026).
    3. Falls back to the default AutoCAD 2021 path if no other is discovered.
    """
    env_path = os.environ.get("ACAD_CONSOLE_PATH")
    if env_path and Path(env_path).exists():
        return env_path

    candidate_paths = [
        r"C:\Program Files\Autodesk\AutoCAD 2021\accoreconsole.exe",
        r"C:\Program Files\Autodesk\AutoCAD 2022\accoreconsole.exe",
        r"C:\Program Files\Autodesk\AutoCAD 2023\accoreconsole.exe",
        r"C:\Program Files\Autodesk\AutoCAD 2024\accoreconsole.exe",
        r"C:\Program Files\Autodesk\AutoCAD 2025\accoreconsole.exe",
        r"C:\Program Files\Autodesk\AutoCAD 2026\accoreconsole.exe",
        r"D:\Program Files\Autodesk\AutoCAD 2021\accoreconsole.exe",
        r"D:\Program Files\Autodesk\AutoCAD 2022\accoreconsole.exe",
        r"D:\Program Files\Autodesk\AutoCAD 2023\accoreconsole.exe",
        r"D:\Program Files\Autodesk\AutoCAD 2024\accoreconsole.exe",
        r"D:\Program Files\Autodesk\AutoCAD 2025\accoreconsole.exe",
    ]
    for cand in candidate_paths:
        if Path(cand).exists():
            return cand

    for base in [Path(r"C:\Program Files\Autodesk"), Path(r"D:\Program Files\Autodesk")]:
        if base.exists() and base.is_dir():
            found = sorted(base.glob("AutoCAD */accoreconsole.exe"), reverse=True)
            if found and found[0].exists():
                return str(found[0])

    return env_path or r"C:\Program Files\Autodesk\AutoCAD 2021\accoreconsole.exe"


ACAD_CONSOLE_PATH = get_accoreconsole_path()

AGENTE_ROOT = Path(
    os.environ.get("AGENTE_ROOT", Path(__file__).resolve().parent.parent)
).resolve()
TEMP_DIR = AGENTE_ROOT / "mcp-autocad" / "temp"
LOG_DIR = AGENTE_ROOT / "mcp-autocad" / "logs"
SKILLS_ROOT = AGENTE_ROOT / ".agents" / "skills"
LISP_LIB_ROOT = AGENTE_ROOT / "lisp-library"

TEMP_DIR.mkdir(parents=True, exist_ok=True)
LOG_DIR.mkdir(parents=True, exist_ok=True)
LISP_LIB_ROOT.mkdir(parents=True, exist_ok=True)


def get_allowed_lisp_roots() -> List[Path]:
    """Returns the list of directories from which AutoLISP (load ...) calls are permitted."""
    roots = [
        SKILLS_ROOT.resolve(),
        LISP_LIB_ROOT.resolve(),
    ]
    extra = os.environ.get("ACAD_EXTRA_LISP_DIRS", "")
    if extra:
        for d in re.split(r"[;,]", extra):
            d = d.strip()
            if d:
                roots.append(Path(d).resolve())
    return roots


ALLOWED_LISP_ROOTS = get_allowed_lisp_roots()

# Logging setup
logging.basicConfig(
    filename=LOG_DIR / "mcp_server.log",
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
)

INSPECTOR_LSP = (
    SKILLS_ROOT
    / "cad-inspector-base"
    / "scripts"
    / "extract_drawing_info.lsp"
)

# Comprehensive circuit breaker forbidden symbols in AutoLISP.
# Evaluated on string-masked code with token boundary regex delimiters
# ((?<![a-zA-Z0-9_\-\*\+\/\=\<\>\!\?\$\%\~\:])...(?![a-zA-Z0-9_\-\*\+\/\=\<\>\!\?\$\%\~\:]))
# to block both direct and indirect calls (e.g. ((cadr '(1 startapp)) ...) or (apply (nth 1 '(a startapp)) ...))
# without false positives on legitimate hyphenated/compound symbols.
FORBIDDEN_LISP_SYMBOLS = [
    # OS execution & processes
    "startapp",
    # File manipulation
    "vl-file-delete",
    "vl-file-rename",
    "vl-file-copy",
    "vl-file-size",
    "vl-file-systime",
    "vl-directory-files",
    "vl-mkdir",
    # Registry manipulation
    "vl-registry-write",
    "vl-registry-delete",
    "vl-registry-read",
    # Environment
    "setenv",
    # Dynamic evaluation & code execution
    "eval",
    "read",
    # ARX / .NET / VBA / ActiveX execution
    "arxload",
    "arxunload",
    "autoarxload",
    "netload",
    "autoload",
    "vl-load-all",
    "vla-eval",
    "vla-runmacro",
    "vla-loaddvb",
    "vla-loadarx",
    "vl-vbaload",
    "vl-vbarun",
    "vla-getinterfaceobject",
    "vlax-import-type-library",
    "vlax-create-object",
    "vlax-get-or-create-object",
    "vlax-invoke",
    "vlax-invoke-method",
    "vla-sendcommand",
    # DOSLib
    r"dos_\w+",
]

LISP_TOKEN_DELIM_LEFT = r"(?<![a-zA-Z0-9_\-\*\+\/\=\<\>\!\?\$\%\~\:])"
LISP_TOKEN_DELIM_RIGHT = r"(?![a-zA-Z0-9_\-\*\+\/\=\<\>\!\?\$\%\~\:])"

DANGEROUS_LISP_CALL_PATTERNS = [
    rf"{LISP_TOKEN_DELIM_LEFT}{sym}{LISP_TOKEN_DELIM_RIGHT}"
    for sym in FORBIDDEN_LISP_SYMBOLS
]

# Patterns that require inspecting string literals (mask_strings=False)
DANGEROUS_LISP_STRING_PATTERNS = [
    r"\bShell\.Application\b",
    r"\bWScript\.Shell\b",
    r"\bScripting\.FileSystemObject\b",
]

DANGEROUS_LISP_PATTERNS = DANGEROUS_LISP_CALL_PATTERNS + DANGEROUS_LISP_STRING_PATTERNS

CONSOLE_COMMAND_ERROR_PATTERNS = [
    r"Unknown command",
    r"Invalid option keyword",
    r"Point or option keyword required",
    r"Requires numeric value",
    r"error:\s+bad argument",
]

LITERAL_LOAD_CALL_RE = re.compile(
    r'^\(\s*load\s+"((?:\\.|[^"\\])+)"(?:\s+(?:nil|t|\'[a-zA-Z0-9_\-]+|"[^"]*"))?\s*\)',
    re.IGNORECASE,
)

mcp = FastMCP("AutoCAD-2021-Agent-MCP")


# ============================================================================
# Security, Encoding & Cleanup Guards
# ============================================================================
def read_lisp_file(path: Path) -> str:
    """Reads an AutoLISP file supporting both UTF-8 (with or without BOM) and cp1252 ANSI."""
    raw = path.read_bytes()
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        return raw.decode("cp1252", errors="replace")


def strip_lisp_comments(lisp_code: str, mask_strings: bool = False) -> str:
    """
    Removes inline ;|...|; and single-line ;... comments outside strings prior to regex scan.
    Preserves exact character offsets by replacing comment characters with spaces.
    When mask_strings=True, also replaces characters inside string literals with spaces.
    """
    out_chars: List[str] = []
    i = 0
    n = len(lisp_code)
    in_str = False
    esc = False
    in_block_comment = False
    in_line_comment = False

    while i < n:
        ch = lisp_code[i]
        nxt = lisp_code[i + 1] if i + 1 < n else ""

        if in_block_comment:
            if ch == "|" and nxt == ";":
                in_block_comment = False
                out_chars.append("  ")
                i += 2
                continue
            out_chars.append(ch if ch in ("\r", "\n") else " ")
            i += 1
            continue

        if in_line_comment:
            if ch in ("\r", "\n"):
                in_line_comment = False
                out_chars.append(ch)
            else:
                out_chars.append(" ")
            i += 1
            continue

        if in_str:
            if esc:
                esc = False
                out_chars.append(ch if (not mask_strings or ch in ("\r", "\n")) else " ")
            elif ch == "\\":
                esc = True
                out_chars.append(ch if not mask_strings else " ")
            elif ch == '"':
                in_str = False
                out_chars.append('"')
            else:
                out_chars.append(ch if (not mask_strings or ch in ("\r", "\n")) else " ")
            i += 1
            continue

        if ch == ";" and nxt == "|":
            in_block_comment = True
            out_chars.append("  ")
            i += 2
            continue
        if ch == ";":
            in_line_comment = True
            out_chars.append(" ")
            i += 1
            continue
        if ch == '"':
            in_str = True
            out_chars.append('"')
            i += 1
            continue

        out_chars.append(ch)
        i += 1

    return "".join(out_chars)


_strip_lisp_comments = strip_lisp_comments


def validate_lisp_structure(lisp_code: str) -> None:
    """
    Validates structural integrity of AutoLISP code:
    - Balanced parentheses '(' and ')'. Depth must never drop below 0 and must finish at 0.
    - Balanced string literals '\"' (taking escaped characters into account).
    - Balanced block comments ';|' and '|;'.
    Prevents premature progn termination or syntax truncation that could accidentally trigger QSAVE.
    """
    paren_depth = 0
    in_str = False
    esc = False
    in_block_comment = False
    in_line_comment = False

    i = 0
    n = len(lisp_code)

    while i < n:
        ch = lisp_code[i]
        nxt = lisp_code[i + 1] if i + 1 < n else ""

        if in_block_comment:
            if ch == "|" and nxt == ";":
                in_block_comment = False
                i += 2
                continue
            i += 1
            continue

        if in_line_comment:
            if ch in ("\r", "\n"):
                in_line_comment = False
            i += 1
            continue

        if in_str:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_str = False
            i += 1
            continue

        if ch == ";" and nxt == "|":
            in_block_comment = True
            i += 2
            continue

        if ch == ";":
            in_line_comment = True
            i += 1
            continue

        if ch == '"':
            in_str = True
            i += 1
            continue

        if ch == "(":
            paren_depth += 1
            i += 1
            continue

        if ch == ")":
            paren_depth -= 1
            if paren_depth < 0:
                raise ValueError("Security circuit breaker: Unbalanced parentheses in AutoLISP code (unexpected ')').")
            i += 1
            continue

        i += 1

    if in_str:
        raise ValueError("Security circuit breaker: Unclosed string literal in AutoLISP code.")
    if in_block_comment:
        raise ValueError("Security circuit breaker: Unclosed block comment ';|...|;' in AutoLISP code.")
    if paren_depth != 0:
        raise ValueError(f"Security circuit breaker: Unbalanced parentheses in AutoLISP code ({paren_depth} unclosed '(').")


FORBIDDEN_COMMAND_NAMES = {
    "SHELL",
    "SH",
    "SCRIPT",
    "NETLOAD",
    "ARX",
    "APPLOAD",
    "VBALOAD",
    "VBARUN",
    "VBAMACRO",
}

SAFE_PROPERTY_COMMANDS = {
    "LAYER",
    "-LAYER",
    "INSERT",
    "-INSERT",
    "STYLE",
    "-STYLE",
    "DIMSTYLE",
    "-DIMSTYLE",
    "TEXT",
    "-TEXT",
    "MTEXT",
    "-MTEXT",
    "CHANGE",
    "-CHANGE",
    "CHPROP",
    "-CHPROP",
    "RENAME",
    "-RENAME",
    "PURGE",
    "-PURGE",
    "COLOR",
    "-COLOR",
    "LINETYPE",
    "-LINETYPE",
    "ATTDEF",
    "-ATTDEF",
    "XREF",
    "-XREF",
    "IMAGE",
    "-IMAGE",
}

LAYER_NAME_OPTIONS = {
    "_SET", "-SET", "SET", "_S", "-S", "S",
    "_MAKE", "-MAKE", "MAKE", "_M", "-M", "M",
    "_NEW", "-NEW", "NEW", "_N", "-N", "N",
    "_RENAME", "-RENAME", "RENAME", "_R", "-R", "R",
    "_FREEZE", "-FREEZE", "FREEZE", "_F", "-F", "F",
    "_THAW", "-THAW", "THAW", "_T", "-T", "T",
    "_LOCK", "-LOCK", "LOCK", "_LO", "-LO", "LO",
    "_UNLOCK", "-UNLOCK", "UNLOCK", "_U", "-U", "U",
    "_ON", "-ON", "ON",
    "_OFF", "-OFF", "OFF",
    "_COLOR", "-COLOR", "COLOR", "_C", "-C", "C",
    "_LTYPE", "-LTYPE", "LTYPE", "_L", "-L", "L",
    "_LWEIGHT", "-LWEIGHT", "LWEIGHT", "_LW", "-LW", "LW",
    "_PLOT", "-PLOT", "PLOT", "_P", "-P", "P",
    "_DESCRIPTION", "-DESCRIPTION", "DESCRIPTION", "_D", "-D", "D",
}

COMMAND_INVOCATION_FORM_RE = re.compile(
    r"\(\s*(?:(?:vl-catch-all-apply|apply)\s+(?:'|\(quote\s+|\(function\s+)?(?:command(?:-s)?|vl-cmdf)\b|(?:command(?:-s)?|vl-cmdf)\b)",
    re.IGNORECASE,
)


def _extract_s_expression(text: str, start_idx: int) -> str:
    depth = 0
    in_s = False
    esc = False
    for i in range(start_idx, len(text)):
        ch = text[i]
        if in_s:
            if esc:
                esc = False
            elif ch == "\\":
                esc = True
            elif ch == '"':
                in_s = False
        else:
            if ch == '"':
                in_s = True
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth == 0:
                    return text[start_idx : i + 1]
    return text[start_idx:]


DYNAMIC_STRING_BUILDERS = {
    "strcat",
    "chr",
    "substr",
    "read",
    "eval",
    "vl-symbol-name",
    "vl-list->string",
    "vl-string-translate",
    "vl-prin1-to-string",
    "vl-princ-to-string",
}


def _is_forbidden_cmd_string(raw_str: str) -> bool:
    clean = re.sub(r"^[_.\-\s]+", "", raw_str).strip().upper()
    clean_unesc = re.sub(r"\\[nNrtP~]", " ", clean)
    first_word = re.split(r"[\s\r\n\t]+", clean_unesc)[0] if clean_unesc else ""
    return first_word in FORBIDDEN_COMMAND_NAMES


def _check_command_shell_invocations(scan_target: str) -> None:
    """
    Checks arguments inside (command ...), (command-s ...), (vl-cmdf ...),
    and (apply 'command ...) / (apply 'vl-cmdf ...) forms for forbidden command invocations.
    Implements 5-rule strict AST & Taint Tracking:
      A. (command) without arguments is allowed (clean cancel like ESC).
      B. First argument in command calls with arguments must be a static string literal.
      C. Dynamic string construction functions (strcat, chr, etc.) inside command forms are forbidden.
      D. Variables assigned forbidden command names via (setq ...) are tracked as tainted and forbidden in command forms.
      E. All string literals inside command forms are scanned against FORBIDDEN_COMMAND_NAMES.
    """
    # 1. Taint tracking for variables assigned forbidden command names (case-insensitive & transitive)
    tainted_vars = set()
    setq_re = re.compile(r"\(\s*setq\b", re.IGNORECASE)
    for m in setq_re.finditer(scan_target):
        form = _extract_s_expression(scan_target, m.start())
        head_m = re.match(r"^\(\s*setq\b", form, re.IGNORECASE)
        if not head_m:
            continue
        body = form[head_m.end() : -1].strip() if form.endswith(")") else form[head_m.end() :].strip()
        tokens = []
        pos = 0
        while pos < len(body):
            while pos < len(body) and body[pos] in " \t\r\n":
                pos += 1
            if pos >= len(body):
                break
            if body[pos] == '"':
                s_end = pos + 1
                esc = False
                while s_end < len(body):
                    if esc:
                        esc = False
                    elif body[s_end] == "\\":
                        esc = True
                    elif body[s_end] == '"':
                        s_end += 1
                        break
                    s_end += 1
                tokens.append(("str", body[pos:s_end]))
                pos = s_end
            elif body[pos] == "(":
                sub_sexpr = _extract_s_expression(body, pos)
                tokens.append(("expr", sub_sexpr))
                pos += len(sub_sexpr)
            else:
                sym_start = pos
                while pos < len(body) and body[pos] not in " \t\r\n()\"":
                    pos += 1
                tokens.append(("sym", body[sym_start:pos]))

        idx = 0
        while idx + 1 < len(tokens):
            var_type, var_name = tokens[idx]
            val_type, val_content = tokens[idx + 1]
            if var_type == "sym":
                is_tainted = False
                if val_type == "str":
                    raw_val = val_content[1:-1] if len(val_content) >= 2 else val_content.strip('"')
                    if _is_forbidden_cmd_string(raw_val):
                        is_tainted = True
                elif val_type == "sym":
                    if val_content.lower() in tainted_vars:
                        is_tainted = True
                elif val_type == "expr":
                    for d_func in DYNAMIC_STRING_BUILDERS:
                        if re.search(rf"\(\s*{re.escape(d_func)}\b", val_content, re.IGNORECASE):
                            is_tainted = True
                            break
                    if not is_tainted:
                        for str_m in re.finditer(r'"((?:\\.|[^"\\])*)"', val_content):
                            if _is_forbidden_cmd_string(str_m.group(1)):
                                is_tainted = True
                                break
                    if not is_tainted:
                        masked_expr = strip_lisp_comments(val_content, mask_strings=True)
                        for t_var in tainted_vars:
                            if re.search(rf"{LISP_TOKEN_DELIM_LEFT}{re.escape(t_var)}{LISP_TOKEN_DELIM_RIGHT}", masked_expr, re.IGNORECASE):
                                is_tainted = True
                                break
                if is_tainted:
                    tainted_vars.add(var_name.lower())
            idx += 2

    # 1b. Block indirect or higher-order references to command / command-s / vl-cmdf (e.g. (mapcar 'command ...))
    masked_all = strip_lisp_comments(scan_target, mask_strings=True)
    valid_cmd_head_spans = [
        (m.start(), m.end()) for m in COMMAND_INVOCATION_FORM_RE.finditer(scan_target)
    ]
    cmd_sym_re = re.compile(
        rf"{LISP_TOKEN_DELIM_LEFT}(?:command(?:-s)?|vl-cmdf){LISP_TOKEN_DELIM_RIGHT}",
        re.IGNORECASE,
    )
    for sym_m in cmd_sym_re.finditer(masked_all):
        s_pos = sym_m.start()
        if not any(h_start <= s_pos <= h_end for h_start, h_end in valid_cmd_head_spans):
            raise ValueError(
                f"Security circuit breaker: Indirect or higher-order reference to '{sym_m.group(0)}' outside direct command invocation is forbidden."
            )

    # 2. Command invocation forms inspection
    for m in COMMAND_INVOCATION_FORM_RE.finditer(scan_target):
        form = _extract_s_expression(scan_target, m.start())
        masked_form = strip_lisp_comments(form, mask_strings=True)

        # Rule C: prohibition of dynamic string builders inside command form
        for d_func in DYNAMIC_STRING_BUILDERS:
            if re.search(rf"\(\s*{re.escape(d_func)}{LISP_TOKEN_DELIM_RIGHT}", masked_form, re.IGNORECASE):
                raise ValueError(
                    f"Security circuit breaker: Dynamic string construction function '({d_func} ...)' is strictly forbidden inside command forms."
                )

        # Rule D: check if any tainted variable is used in form (using exact AutoLISP token boundaries)
        for t_var in tainted_vars:
            if re.search(rf"{LISP_TOKEN_DELIM_LEFT}{re.escape(t_var)}{LISP_TOKEN_DELIM_RIGHT}", masked_form, re.IGNORECASE):
                raise ValueError(
                    f"Security circuit breaker: Tainted variable '{t_var}' containing forbidden command name used in command invocation."
                )

        # Rule A & B: check direct command calls and apply calls
        direct_match = re.match(r"^\(\s*(command(?:-s)?|vl-cmdf)\b", form, re.IGNORECASE)
        if direct_match:
            args_body = form[direct_match.end() : -1].strip() if form.endswith(")") else form[direct_match.end() :].strip()
            # Rule A: 0 arguments allowed (command cancellation)
            if not args_body:
                continue

            pos = 0
            while pos < len(args_body) and args_body[pos] in " \t\r\n":
                pos += 1
            if pos < len(args_body):
                if args_body[pos] == '"':
                    pass  # Rule B: first argument is string literal
                elif args_body[pos] == "(":
                    raise ValueError(
                        f"Security circuit breaker: Dynamic or variable command names are strictly forbidden in command/vl-cmdf: '{form}'"
                    )
                else:
                    sym_end = pos
                    while sym_end < len(args_body) and args_body[sym_end] not in " \t\r\n()\"":
                        sym_end += 1
                    var_name = args_body[pos:sym_end]
                    raise ValueError(
                        f"Security circuit breaker: Variable command name '{var_name}' is strictly forbidden in command/vl-cmdf."
                    )
        else:
            apply_match = re.match(
                r"^\(\s*(?:vl-catch-all-apply|apply)\s+(?:'|\(\s*quote\s+|\(\s*function\s+)?(?:command(?:-s)?|vl-cmdf)\b\s*\)?",
                form,
                re.IGNORECASE,
            )
            if apply_match:
                rest_apply = form[apply_match.end() : -1].strip() if form.endswith(")") else form[apply_match.end() :].strip()
                if not re.match(r"^(?:'\(\s*\"|\(\s*quote\s+\(\s*\"|\(\s*list\s+\")", rest_apply, re.IGNORECASE):
                    raise ValueError(
                        f"Security circuit breaker: Dynamic argument list in apply command/vl-cmdf is forbidden: '{form}'"
                    )

        # Rule E: inspect string literals in form against FORBIDDEN_COMMAND_NAMES
        str_matches = list(re.finditer(r'"((?:\\.|[^"\\])*)"', form))
        if str_matches:
            primary_cmd_raw = str_matches[0].group(1)
            # Primary command cannot be forbidden
            if _is_forbidden_cmd_string(primary_cmd_raw):
                raise ValueError(
                    f"Security circuit breaker: Detected forbidden command invocation '{primary_cmd_raw}' in AutoLISP."
                )

            primary_clean = re.sub(r"^[_.\-\s]+", "", primary_cmd_raw).strip().upper()
            is_property_cmd = primary_clean in SAFE_PROPERTY_COMMANDS

            for idx in range(1, len(str_matches)):
                arg_str = str_matches[idx].group(1)
                if not _is_forbidden_cmd_string(arg_str):
                    continue

                # arg_str matches a forbidden command name (e.g. SHELL, SH, SCRIPT, etc.)
                # Check if it can be legitimately allowed strictly as a CAD entity/property parameter:
                # 1. It must NOT start with command prefix (_ . -)
                has_cmd_prefix = bool(re.match(r"^[_.\-]", arg_str.strip()))
                # 2. It must NOT contain arguments/whitespace/newlines (e.g. "SHELL calc")
                has_internal_args = bool(re.search(r"[\s\r\n\t\\]", arg_str.strip()))
                # 3. No command terminator ("") can have appeared prior to this token.
                #    In AutoCAD, "" sends Enter and terminates the active command, returning to the Command: prompt.
                #    Any token after "" is evaluated at the Command: prompt and would execute as a new command!
                terminated_before = any(str_matches[k].group(1) == "" for k in range(idx))

                is_safe_param = False
                allow_selection_terminator = (
                    primary_clean in ("CHPROP", "-CHPROP", "CHANGE", "-CHANGE")
                    and sum(1 for k in range(idx) if str_matches[k].group(1) == "") <= 1
                )
                if is_property_cmd and not has_cmd_prefix and not has_internal_args and (not terminated_before or allow_selection_terminator):
                    prev_str = str_matches[idx - 1].group(1).strip().upper() if idx > 0 else ""
                    prev2_str = str_matches[idx - 2].group(1).strip().upper() if idx > 1 else ""
                    next_str = str_matches[idx + 1].group(1).strip().upper() if idx + 1 < len(str_matches) else ""
                    layer_val_opts = {
                        "_COLOR", "-COLOR", "COLOR", "_C", "-C", "C",
                        "_LTYPE", "-LTYPE", "LTYPE", "_L", "-L", "L",
                        "_LWEIGHT", "-LWEIGHT", "LWEIGHT", "_LW", "-LW", "LW",
                    }
                    if primary_clean in ("LAYER", "-LAYER"):
                        # Layer name immediately follows a layer option, or follows <option> <val> (e.g. _Color "1" "SHELL")
                        if prev_str in LAYER_NAME_OPTIONS or (prev2_str in layer_val_opts and prev_str not in LAYER_NAME_OPTIONS and prev_str != ""):
                            if next_str == "" or next_str in LAYER_NAME_OPTIONS:
                                is_safe_param = True
                    elif primary_clean in ("INSERT", "-INSERT"):
                        if idx == 1:
                            is_safe_param = True
                    elif primary_clean in ("STYLE", "-STYLE"):
                        if idx == 1:
                            is_safe_param = True
                    elif primary_clean in ("TEXT", "-TEXT"):
                        # In TEXT, the text string is the final argument of the command; no extra strings or tokens may follow it
                        trailing_after_str = form[str_matches[idx].end() : -1].strip() if form.endswith(")") else form[str_matches[idx].end() :].strip()
                        if idx == len(str_matches) - 1 and trailing_after_str == "":
                            if idx == 1 or (idx == 2 and prev_str in ("_J", "-J", "J", "_S", "-S", "S", "_C", "_M", "_R", "_TL", "_TC", "_TR", "_ML", "_MC", "_MR", "_BL", "_BC", "_BR")):
                                is_safe_param = True
                    elif primary_clean in ("MTEXT", "-MTEXT"):
                        # In MTEXT, the text string must be the last string or followed only by ""
                        if idx == len(str_matches) - 1 or (idx == len(str_matches) - 2 and next_str == ""):
                            is_safe_param = True
                    elif primary_clean in ("PURGE", "-PURGE"):
                        purge_cats = {"_LA", "-LA", "LA", "_LAYERS", "_B", "-B", "B", "_BLOCKS", "_ST", "_D", "_LT", "_MA", "_SH", "_TR"}
                        if idx == 2 and prev_str in purge_cats and next_str in ("_N", "N", "_Y", "Y", ""):
                            is_safe_param = True
                    elif primary_clean in ("RENAME", "-RENAME"):
                        rename_cats = {"_LA", "-LA", "LA", "_B", "-B", "B", "_ST", "_D", "_LT", "_U", "_VI", "_VP"}
                        if (idx == 2 and prev_str in rename_cats) or (idx == 3 and prev2_str in rename_cats):
                            is_safe_param = True
                    elif primary_clean in ("CHPROP", "-CHPROP", "CHANGE", "-CHANGE"):
                        if prev_str in ("_LA", "-LA", "LA", "_LAYER", "_LT", "_LTYPE", "_C", "_COLOR") and (next_str == "" or next_str.startswith("_")):
                            is_safe_param = True
                    elif primary_clean in ("DIMSTYLE", "-DIMSTYLE"):
                        if idx == 1 or prev_str in ("_R", "_RESTORE", "R", "_S", "_SAVE", "S"):
                            is_safe_param = True

                if is_safe_param:
                    continue

                raise ValueError(
                    f"Security circuit breaker: Detected forbidden command invocation '{arg_str}' in AutoLISP."
                )


def _check_open_calls_whitelist(scan_target: str, masked_target: str) -> None:
    """
    Strict whitelist policy: Any appearance of the token 'open' in AutoLISP code
    must be a direct function call (open <filepath-expr> "r") with literal read-only mode "r" or "R".
    Any write/append mode, indirect variable usage, or dynamic mode variable is strictly rejected.
    """
    open_token_re = re.compile(rf"{LISP_TOKEN_DELIM_LEFT}open{LISP_TOKEN_DELIM_RIGHT}", re.IGNORECASE)
    for match in open_token_re.finditer(masked_target):
        start_open = match.start()
        idx = start_open - 1
        while idx >= 0 and masked_target[idx].isspace():
            idx -= 1
        if idx < 0 or masked_target[idx] != "(":
            raise ValueError(
                "Security circuit breaker: Indirect, quoted, or variable usage of symbol 'open' is forbidden."
            )

        i = match.end()
        n = len(scan_target)
        while i < n and scan_target[i].isspace():
            i += 1
        if i >= n:
            raise ValueError("Security circuit breaker: Malformed 'open' call (missing arguments).")

        # Arg 1: <expr> (string literal, balanced S-expression, or atom)
        if scan_target[i] == '"':
            i += 1
            esc = False
            while i < n:
                ch = scan_target[i]
                if esc:
                    esc = False
                elif ch == "\\":
                    esc = True
                elif ch == '"':
                    i += 1
                    break
                i += 1
        elif scan_target[i] == "(":
            depth = 0
            in_s = False
            esc = False
            while i < n:
                ch = scan_target[i]
                if in_s:
                    if esc:
                        esc = False
                    elif ch == "\\":
                        esc = True
                    elif ch == '"':
                        in_s = False
                else:
                    if ch == '"':
                        in_s = True
                    elif ch == "(":
                        depth += 1
                    elif ch == ")":
                        depth -= 1
                        if depth == 0:
                            i += 1
                            break
                i += 1
        else:
            while i < n and not scan_target[i].isspace() and scan_target[i] not in "()":
                i += 1

        while i < n and scan_target[i].isspace():
            i += 1
        if i >= n:
            raise ValueError("Security circuit breaker: Malformed 'open' call (missing mode argument).")

        mode_match = re.match(r'^"(?:r|R)"', scan_target[i:], re.IGNORECASE)
        if not mode_match:
            raise ValueError(
                "Security circuit breaker: AutoLISP 'open' is only permitted in read-only mode (\"r\")."
            )

        i += mode_match.end()
        while i < n and scan_target[i].isspace():
            i += 1
        if i >= n or scan_target[i] != ")":
            raise ValueError(
                "Security circuit breaker: Malformed 'open' call (unexpected extra arguments or unclosed form)."
            )


def _resolve_and_verify_skill_load_path(raw_target: str) -> Path:
    """Resolves a literal path from (load "...") and ensures it is an existing .lsp file inside ALLOWED_LISP_ROOTS."""
    unescaped = raw_target.replace("\\\\", "\\")
    candidate = Path(unescaped)
    if candidate.suffix == "":
        candidate = candidate.with_suffix(".lsp")

    if not candidate.is_absolute():
        if not (AGENTE_ROOT / candidate).exists() and (LISP_LIB_ROOT / candidate).exists():
            candidate = (LISP_LIB_ROOT / candidate).resolve()
        elif (SKILLS_ROOT / candidate).exists():
            candidate = (SKILLS_ROOT / candidate).resolve()
        else:
            slug = re.sub(r"[^a-z0-9-]+", "-", candidate.stem.lower()).strip("-")
            skill_script_candidate = SKILLS_ROOT / slug / "scripts" / f"{slug}.lsp"
            if skill_script_candidate.exists():
                candidate = skill_script_candidate.resolve()
            else:
                candidate = (AGENTE_ROOT / candidate).resolve()
    else:
        candidate = candidate.resolve()

    if candidate.suffix.lower() != ".lsp":
        raise ValueError(
            f"Security circuit breaker: (load '{raw_target}') must target a .lsp file."
        )

    if not candidate.exists():
        raise FileNotFoundError(
            f"Security circuit breaker: LISP file '{candidate}' does not exist."
        )

    allowed_roots = get_allowed_lisp_roots()
    is_allowed = False
    for root in allowed_roots:
        try:
            candidate.relative_to(root)
            is_allowed = True
            break
        except ValueError:
            pass

    if not is_allowed:
        raise ValueError(
            f"Security circuit breaker: (load '{raw_target}') is outside trusted directories ({[str(r) for r in allowed_roots]})."
        )

    return candidate

def validate_lisp_safety(
    lisp_code: str,
    _depth: int = 0,
    allow_trusted_file_write: bool = False,
) -> None:
    """
    Circuit breaker: ensures no dangerous OS-level functions are executed via AutoLISP.
    CAD routines must only manipulate drawing entities and tables.
    Permits literal (load "...") calls ONLY when targeting existing .lsp files inside ALLOWED_LISP_ROOTS.
    """
    if _depth > 5:
        raise ValueError("Security circuit breaker: Maximum (load ...) nesting depth exceeded.")

    scan_target = strip_lisp_comments(lisp_code, mask_strings=False)
    masked_target = strip_lisp_comments(lisp_code, mask_strings=True)

    for pattern in DANGEROUS_LISP_CALL_PATTERNS:
        match = re.search(pattern, masked_target, re.IGNORECASE)
        if match:
            raise ValueError(
                f"Security circuit breaker: Detected forbidden OS/system call '{match.group(0).strip()}' in AutoLISP routine."
            )

    for pattern in DANGEROUS_LISP_STRING_PATTERNS:
        match = re.search(pattern, scan_target, re.IGNORECASE)
        if match:
            raise ValueError(
                f"Security circuit breaker: Detected forbidden OS/system call '{match.group(0)}' in AutoLISP routine."
            )

    _check_command_shell_invocations(scan_target)

    if not allow_trusted_file_write:
        _check_open_calls_whitelist(scan_target, masked_target)

    for load_match in re.finditer(r"\(\s*load\b", masked_target, re.IGNORECASE):
        pos = load_match.start()
        call_match = LITERAL_LOAD_CALL_RE.match(scan_target[pos:])
        if not call_match:
            raise ValueError(
                "Security circuit breaker: Dynamic or non-literal (load ...) calls are forbidden."
            )
        raw_target = call_match.group(1)
        candidate = _resolve_and_verify_skill_load_path(raw_target)
        if candidate.exists() and candidate.is_file():
            validate_lisp_safety(
                read_lisp_file(candidate),
                _depth=_depth + 1,
                allow_trusted_file_write=False,
            )


def normalize_skill_loads(lisp_code: str) -> str:
    """
    Rewrites validated (load "...") calls outside strings/comments to canonical POSIX paths inside ALLOWED_LISP_ROOTS
    so AutoCAD resolves them reliably regardless of the active DWG directory.
    """
    scan_target = strip_lisp_comments(lisp_code, mask_strings=False)
    masked_target = strip_lisp_comments(lisp_code, mask_strings=True)
    replacements: List[Tuple[int, int, str]] = []
    for load_match in re.finditer(r"\(\s*load\b", masked_target, re.IGNORECASE):
        pos = load_match.start()
        call_match = LITERAL_LOAD_CALL_RE.match(scan_target[pos:])
        if call_match:
            raw_target = call_match.group(1)
            candidate = _resolve_and_verify_skill_load_path(raw_target)
            abs_start = pos + call_match.start(1)
            abs_end = pos + call_match.end(1)
            replacements.append((abs_start, abs_end, candidate.as_posix()))
    out = lisp_code
    for start_idx, end_idx, new_path in reversed(replacements):
        out = out[:start_idx] + new_path + out[end_idx:]
    return out


def resolve_skill_script_code(skill_script: Optional[str]) -> str:
    """
    Safely loads an approved skill .lsp script from SKILLS_ROOT or LISP_LIB_ROOT by slug or relative path,
    preventing arbitrary filesystem traversal while allowing execution without hitting the (load) blocker.
    """
    if not skill_script:
        return ""
    candidate = Path(skill_script)
    if not candidate.is_absolute():
        # Check direct existence in LISP_LIB_ROOT, SKILLS_ROOT or AGENTE_ROOT
        if (LISP_LIB_ROOT / candidate).exists():
            candidate = (LISP_LIB_ROOT / candidate).resolve()
        elif (SKILLS_ROOT / candidate).exists():
            candidate = (SKILLS_ROOT / candidate).resolve()
        elif (AGENTE_ROOT / candidate).exists():
            candidate = (AGENTE_ROOT / candidate).resolve()
        else:
            slug = re.sub(r"[^a-z0-9-]+", "-", candidate.stem.lower()).strip("-")
            scripts_dir = SKILLS_ROOT / slug / "scripts"
            candidate_file = scripts_dir / f"{slug}.lsp"
            if not candidate_file.exists() and scripts_dir.is_dir():
                lsp_files = sorted(scripts_dir.glob("*.lsp"))
                if len(lsp_files) == 1:
                    candidate_file = lsp_files[0]
            if not candidate_file.exists():
                lib_candidate = LISP_LIB_ROOT / f"{slug}.lsp"
                if lib_candidate.exists():
                    candidate_file = lib_candidate
            if candidate_file.exists():
                candidate = candidate_file.resolve()
            else:
                candidate = (AGENTE_ROOT / candidate).resolve()
    else:
        candidate = candidate.resolve()

    allowed_roots = get_allowed_lisp_roots()
    is_allowed = False
    for root in allowed_roots:
        try:
            candidate.relative_to(root)
            is_allowed = True
            break
        except ValueError:
            pass

    if not is_allowed:
        raise ValueError(
            f"Security error: skill_script '{skill_script}' must reside inside trusted directories ({[str(r) for r in allowed_roots]})."
        )

    if not candidate.exists():
        raise FileNotFoundError(f"Skill AutoLISP script not found: {candidate}")

    code = read_lisp_file(candidate)
    validate_lisp_structure(code)
    validate_lisp_safety(code, allow_trusted_file_write=False)
    return code


def cleanup_stale_temp_files(prefix: Optional[str] = None, max_age_sec: float = 300.0) -> None:
    """Removes temporary .lsp, .scr, .json, and .status files older than max_age_sec so deferred live loads are not unlinked early."""
    now = time.time()
    try:
        for item in TEMP_DIR.iterdir():
            if item.name == ".gitkeep":
                continue
            if prefix and not item.name.startswith(prefix):
                continue
            if item.is_file() and (now - item.stat().st_mtime) > max_age_sec:
                item.unlink(missing_ok=True)
    except Exception:
        pass


def decode_accore_stdout(raw: bytes) -> str:
    """
    Resilient decoding of accoreconsole.exe stdout bytes.
    Handles UTF-16LE, odd-byte misalignments caused by ARX/CRX modules, and CP1252 ANSI fallback.
    """
    if not raw:
        return ""
    text_u16 = raw.decode("utf-16le", errors="replace")
    if "[MCP_" in text_u16 or "AutoCAD" in text_u16:
        return text_u16
    if b"\x00" in raw:
        aligned = raw[1:].decode("utf-16le", errors="replace")
        if "[MCP_" in aligned or "AutoCAD" in aligned:
            return aligned
        return raw.replace(b"\x00", b"").decode("cp1252", errors="replace")
    return raw.decode("cp1252", errors="replace")


def detect_console_command_warnings(console_text: str) -> List[str]:
    """Detects command-line warnings or bad argument messages emitted by AutoCAD strictly within mutation execution window."""
    start_marker = "[MCP_MUTATE_EXEC_START]"
    end_marker = "[MCP_MUTATE_EXEC_END]"
    if start_marker in console_text:
        start_idx = console_text.find(start_marker) + len(start_marker)
        end_idx = console_text.find(end_marker, start_idx)
        scan_window = console_text[start_idx:end_idx] if end_idx != -1 else console_text[start_idx:]
    else:
        scan_window = console_text

    warnings = []
    for pattern in CONSOLE_COMMAND_ERROR_PATTERNS:
        for match in re.finditer(pattern, scan_window, re.IGNORECASE):
            warnings.append(match.group(0))
    return warnings


def wrap_untrusted_cad_text(raw_text: Any, tag_type: str = "drawing_text") -> str:
    """Wraps text extracted from CAD drawings in strict XML delimiters with HTML escaping."""
    if raw_text is None:
        return ""
    text_str = str(raw_text)
    escaped = html.escape(text_str, quote=True)
    return f"<untrusted_cad_data type='{tag_type}'>{escaped}</untrusted_cad_data>"


def safe_com_getattr(obj: Any, attr: str, default: Any = None) -> Any:
    """Safely retrieves a COM property without raising exceptions on unsupported interfaces."""
    try:
        return getattr(obj, attr)
    except Exception:
        pass
    return default


# ============================================================================
# Lock & File Helpers (.dwl / .dwl2 Inspection & Sequential Legacy Backups)
# ============================================================================
def get_dwg_lock_owner_info(dwg_path: Path) -> Optional[Dict[str, str]]:
    """Reads user, machine, and datetime from AutoCAD lock files (.dwl / .dwl2)."""
    dwl2_path = dwg_path.with_suffix(".dwl2")
    dwl_path = dwg_path.with_suffix(".dwl")

    user, machine, dt = "Usuario desconocido", "Equipo desconocido", "Hora desconocida"
    found_dwl2 = False

    if dwl2_path.exists():
        try:
            tree = ET.parse(dwl2_path)
            root = tree.getroot()
            tag_map = {elem.tag.lower(): (elem.text or "").strip() for elem in root.iter()}
            u = tag_map.get("username", "")
            m = tag_map.get("machinename", "")
            d = tag_map.get("datetime", "")
            if u:
                user = u
            if m:
                machine = m
            if d:
                dt = d
            if user != "Usuario desconocido" and machine != "Equipo desconocido":
                return {"user": user, "machine": machine, "datetime": dt}
            found_dwl2 = True
        except Exception:
            pass

    if dwl_path.exists():
        try:
            lines = [
                line.strip()
                for line in dwl_path.read_text(encoding="latin-1", errors="replace").splitlines()
                if line.strip()
            ]
            if len(lines) >= 1 and lines[0]:
                user = lines[0]
            if len(lines) >= 2 and lines[1]:
                machine = lines[1]
            if len(lines) >= 3 and lines[2]:
                dt = lines[2]
            return {"user": user, "machine": machine, "datetime": dt}
        except Exception:
            pass

    if found_dwl2:
        return {"user": user, "machine": machine, "datetime": dt}

    return None


def check_dwg_file_status(dwg_path: Path) -> Tuple[bool, Optional[str], Optional[Dict[str, Any]]]:
    """
    Checks DWG accessibility:
    Distinguishes Windows filesystem Read-Only attributes, ACL permission denials, and active AutoCAD locks.
    Returns (is_blocked, user_friendly_message, lock_owner_info).
    """
    if not dwg_path.exists():
        return False, None, None

    # 1. Check Windows Read-Only filesystem attribute
    file_stat = dwg_path.stat()
    if not (file_stat.st_mode & stat.S_IWRITE):
        owner = get_dwg_lock_owner_info(dwg_path)
        msg = (
            f"El archivo '{dwg_path.name}' tiene atributo de 'Solo Lectura' en el sistema de archivos de Windows. "
            "Quitale el atributo de solo lectura antes de intentar editarlo."
        )
        return True, msg, owner

    # 2. Check Win32 kernel exclusive lock (dwShareMode=0) vs ACL access denied
    is_kernel_locked = False
    access_denied = False
    try:
        handle = win32file.CreateFile(
            str(dwg_path),
            win32file.GENERIC_READ | win32file.GENERIC_WRITE,
            0,  # 0 = Exclusive share mode; fails if any process holds an open handle
            None,
            win32file.OPEN_EXISTING,
            win32file.FILE_ATTRIBUTE_NORMAL,
            None,
        )
        win32file.CloseHandle(handle)
    except pywintypes.error as exc:
        winerror = getattr(exc, "winerror", getattr(exc, "args", [None])[0])
        if winerror in (32, 33):  # ERROR_SHARING_VIOLATION, ERROR_LOCK_VIOLATION
            is_kernel_locked = True
        elif winerror == 5:  # ERROR_ACCESS_DENIED
            access_denied = True
        else:
            is_kernel_locked = True
    except OSError as exc:
        winerror = getattr(exc, "winerror", None)
        if winerror in (32, 33):
            is_kernel_locked = True
        elif winerror == 5:
            access_denied = True
        else:
            is_kernel_locked = True

    owner = get_dwg_lock_owner_info(dwg_path)
    local_pc = os.environ.get("COMPUTERNAME", "").strip().upper()

    if access_denied:
        msg = (
            f"Acceso denegado (permisos ACL o de red insuficientes) al intentar acceder a '{dwg_path.name}'. "
            "Verifica los permisos de usuario en la carpeta o recurso compartido."
        )
        return True, msg, owner

    if is_kernel_locked:
        owner_kernel: Dict[str, Any] = {**(owner or {}), "is_kernel_locked": True}
        if owner and owner.get("user") != "Usuario desconocido":
            msg = (
                f"El archivo '{dwg_path.name}' está actualmente abierto por el usuario '{owner['user']}' "
                f"en el equipo '{owner['machine']}' (desde {owner['datetime']}). "
                "Cerralo para editarlo en modo Headless, o pedime que use el modo Live si es tu ventana activa."
            )
        else:
            msg = (
                f"El archivo '{dwg_path.name}' está actualmente bloqueado por un proceso de AutoCAD abierto. "
                "Cerrá el archivo o usá 'execute_cad_live'."
            )
        return True, msg, owner_kernel

    # 3. Kernel handle is free locally; check if another workstation holds a remote .dwl/.dwl2 lock
    if owner:
        lock_pc = owner.get("machine", "").strip().upper()
        if lock_pc and lock_pc != "EQUIPO DESCONOCIDO" and local_pc and lock_pc != local_pc:
            dwl_path = dwg_path.with_suffix(".dwl")
            dwl2_path = dwg_path.with_suffix(".dwl2")
            mtime = 0.0
            for p in (dwl_path, dwl2_path):
                if p.exists():
                    mtime = max(mtime, p.stat().st_mtime)
            age_hours = (time.time() - mtime) / 3600.0 if mtime > 0 else 0.0
            is_stale = age_hours > 12.0

            stale_note = (
                f" (Aviso: el bloqueo tiene {age_hours:.1f} horas de antigüedad y podría ser un bloqueo huérfano tras un cierre inesperado)."
                if is_stale
                else ""
            )
            msg = (
                f"El archivo '{dwg_path.name}' registra un bloqueo activo de otra terminal ('{lock_pc}', usuario '{owner['user']}'){stale_note}."
            )
            owner_data: Dict[str, Any] = {
                **owner,
                "is_kernel_locked": False,
                "is_stale": is_stale,
                "age_hours": round(age_hours, 1),
            }
            return True, msg, owner_data

    return False, None, None


def is_dwg_locked(dwg_path: Path) -> bool:
    """Helper returning simple boolean lock state."""
    blocked, _, _ = check_dwg_file_status(dwg_path)
    return blocked


def create_sequential_legacy_backup(dwg_path: Path) -> Path:
    """Creates a sequential backup file: <OriginalName>-Legacy-0.dwg, -Legacy-1.dwg, etc."""
    base_stem = dwg_path.stem
    i = 0
    while True:
        candidate = dwg_path.with_name(f"{base_stem}-Legacy-{i}.dwg")
        if not candidate.exists():
            shutil.copy2(dwg_path, candidate)
            return candidate
        i += 1


def safe_restore_backup(backup_path: Path, target_dwg: Path, remove_redundant_backup: bool = True) -> bool:
    """
    Restores backup_path back to target_dwg with retries against filesystem transient locks.
    If remove_redundant_backup=True, atomically replaces target_dwg via os.replace so no redundant -Legacy-N file remains.
    """
    for attempt in range(5):
        try:
            if target_dwg.exists():
                os.chmod(str(target_dwg), stat.S_IWRITE)
            if remove_redundant_backup:
                os.replace(str(backup_path), str(target_dwg))
            else:
                shutil.copy2(str(backup_path), str(target_dwg))
            return True
        except Exception:
            time.sleep(0.3 * (attempt + 1))
    logging.error(f"Failed to restore backup {backup_path} to {target_dwg} after 5 attempts.")
    return False


def is_development_staging_dir(root: Optional[Path] = None) -> bool:
    """
    Checks whether AGENTE_ROOT is the development/debugging staging folder (e.g. Laboral/CAD Proyect),
    where the MCP server and global IDE skill junctions must NOT be auto-installed.
    """
    if os.environ.get("ACAD_MCP_DEV_NO_INSTALL", "").strip() == "1":
        return True
    if os.environ.get("ACAD_MCP_FORCE_GLOBAL_SYNC", "").strip() == "1":
        return False
    check_root = (root or AGENTE_ROOT).resolve()
    if "cad proyect" in str(check_root).lower():
        return True
    return False


def sync_skill_to_global_ides(skill_dir: Path, clean_slug: str, force: bool = False) -> List[str]:
    """Links or copies a skill folder into global Antigravity and Codex skill directories (skipped in development staging folder)."""
    if not force and is_development_staging_dir():
        logging.info(f"Skipping global IDE skill sync for '{clean_slug}' (development staging mode active).")
        return []
    synced_locations: List[str] = []
    home = Path.home()
    target_dirs = [
        home / ".gemini" / "config" / "skills" / clean_slug,
        home / ".codex" / "skills" / clean_slug,
    ]
    for target in target_dirs:
        try:
            target.parent.mkdir(parents=True, exist_ok=True)
            if not target.exists():
                try:
                    subprocess.run(
                        ["cmd", "/c", "mklink", "/J", str(target), str(skill_dir)],
                        check=True,
                        capture_output=True,
                        text=True,
                    )
                except Exception:
                    shutil.copytree(skill_dir, target, dirs_exist_ok=True)
            if not target.exists() or (target.is_dir() and not target.is_symlink() and target.resolve() != skill_dir.resolve()):
                shutil.copytree(skill_dir, target, dirs_exist_ok=True)
            synced_locations.append(str(target))
        except Exception as exc:
            logging.warning(f"Could not sync global skill to {target}: {exc}")
    return synced_locations


def sync_all_skills_to_global_ides() -> List[str]:
    """Iterates through all skill directories in SKILLS_ROOT and syncs them to global IDE dirs."""
    all_synced: List[str] = []
    if SKILLS_ROOT.exists() and SKILLS_ROOT.is_dir():
        for skill_dir in SKILLS_ROOT.iterdir():
            if skill_dir.is_dir() and (skill_dir / "SKILL.md").exists():
                all_synced.extend(sync_skill_to_global_ides(skill_dir, skill_dir.name))
    return all_synced


# ============================================================================
# Git Synchronization (Safe Fast-Forward & Human-Confirmed Commits)
# ============================================================================
LAST_GIT_SYNC_STATUS = "Initializing..."


def git_auto_pull() -> str:
    """Performs a safe fast-forward pull on startup without breaking on divergence."""
    try:
        if not (AGENTE_ROOT / ".git").exists():
            return "Git repository not initialized."
        env = {
            **os.environ,
            "GIT_TERMINAL_PROMPT": "0",
            "GCM_INTERACTIVE": "never",
        }
        res = subprocess.run(
            ["git", "pull", "--ff-only"],
            cwd=str(AGENTE_ROOT),
            capture_output=True,
            text=True,
            env=env,
            timeout=15,
        )
        if res.returncode == 0:
            out = res.stdout.strip() or "Up to date."
            logging.info(f"Git pull successful: {out}")
            return out
        else:
            err = res.stderr.strip()
            logging.warning(f"Git pull --ff-only failed: {err}")
            return f"Git pull warning (manual sync may be required): {err}"
    except Exception as exc:
        logging.error(f"Git pull exception: {exc}")
        return f"Git pull skipped: {exc}"


def git_auto_push(commit_message: str) -> str:
    """Commits, rebases against remote updates, and pushes curated memory and skills to GitHub."""
    try:
        if not (AGENTE_ROOT / ".git").exists():
            return "Saved locally (Git repo not initialized)."
        env = {
            **os.environ,
            "GIT_TERMINAL_PROMPT": "0",
            "GCM_INTERACTIVE": "never",
        }
        subprocess.run(
            ["git", "add", "MEMORY.md", ".agents/", "lisp-library/"],
            cwd=str(AGENTE_ROOT),
            check=True,
            env=env,
            timeout=10,
        )
        commit_res = subprocess.run(
            ["git", "commit", "-m", commit_message],
            cwd=str(AGENTE_ROOT),
            capture_output=True,
            text=True,
            env=env,
            timeout=10,
        )
        combined_out = (commit_res.stdout + "\n" + commit_res.stderr).strip()
        if "nothing to commit" in combined_out.lower():
            return "No new changes to commit."
        if commit_res.returncode != 0:
            logging.error(f"Git commit failed: {combined_out}")
            return f"Saved locally, but Git commit failed: {combined_out}"

        # Pull with rebase prior to push so Home PC and Office PC commits merge cleanly
        pull_res = subprocess.run(
            ["git", "pull", "--rebase", "--autostash"],
            cwd=str(AGENTE_ROOT),
            capture_output=True,
            text=True,
            env=env,
            timeout=20,
        )
        if pull_res.returncode != 0:
            err = (pull_res.stderr or pull_res.stdout).strip()
            subprocess.run(
                ["git", "rebase", "--abort"],
                cwd=str(AGENTE_ROOT),
                capture_output=True,
                text=True,
                env=env,
                timeout=10,
            )
            logging.error(f"Git pull --rebase conflict aborted cleanly: {err}")
            return f"Committed locally, but Git pull --rebase failed (rebase aborted cleanly): {err}"

        push_res = subprocess.run(
            ["git", "push"],
            cwd=str(AGENTE_ROOT),
            capture_output=True,
            text=True,
            env=env,
            timeout=20,
        )
        if push_res.returncode != 0:
            err = (push_res.stderr or push_res.stdout).strip()
            logging.error(f"Git push failed: {err}")
            return f"Committed locally, but Git push failed: {err}"

        logging.info(f"Git push success: {commit_message}")
        return "Successfully committed and pushed to GitHub."
    except Exception as exc:
        logging.error(f"Git push failed: {exc}")
        return f"Saved locally, but Git push failed: {exc}"


def _background_git_pull():
    global LAST_GIT_SYNC_STATUS
    if is_development_staging_dir():
        LAST_GIT_SYNC_STATUS = "Development staging mode (global install & auto-pull skipped)."
        return
    LAST_GIT_SYNC_STATUS = git_auto_pull()
    try:
        sync_all_skills_to_global_ides()
    except Exception as exc:
        logging.warning(f"Startup skill sync failed: {exc}")


# ============================================================================
# COM Helpers for Secondary Live Mode (STA Thread-Safe with Concurrency Lock)
# ============================================================================
COM_THREAD_LOCK = threading.Lock()


def _is_com_call_rejected(exc: Exception) -> bool:
    """Checks whether the exception represents RPC_E_CALL_REJECTED (-2147418111 / 0x80010001)."""
    hresult = getattr(exc, "hresult", None)
    if hresult in (-2147418111, 0x80010001):
        return True
    args = getattr(exc, "args", ())
    if args and isinstance(args[0], int) and args[0] in (-2147418111, 0x80010001):
        return True
    return False


def _get_acad_gui_app(create_if_missing: bool = False) -> Any:
    """Gets the active AutoCAD GUI instance via COM, or optionally creates a new visible instance."""
    try:
        return win32com.client.GetActiveObject(ACAD_PROG_ID)
    except Exception as exc:
        if _is_com_call_rejected(exc):
            raise
    try:
        return win32com.client.GetActiveObject(ACAD_GENERIC_PROG_ID)
    except Exception as exc:
        if _is_com_call_rejected(exc):
            raise
    if create_if_missing:
        try:
            acad = win32com.client.Dispatch(ACAD_PROG_ID)
        except Exception as exc:
            if _is_com_call_rejected(exc):
                raise
            acad = win32com.client.Dispatch(ACAD_GENERIC_PROG_ID)
        try:
            acad.Visible = True
        except Exception:
            pass
        return acad
    raise RuntimeError("AutoCAD GUI application is not currently running.")


def com_retry(max_attempts: int = 5, delay: float = 0.5):
    """Decorator retrying COM calls when AutoCAD rejects them with RPC_E_CALL_REJECTED."""
    def decorator(func):
        @wraps(func)
        def wrapper(*args, **kwargs):
            last_err = None
            for attempt in range(max_attempts):
                try:
                    return func(*args, **kwargs)
                except Exception as exc:
                    if _is_com_call_rejected(exc):
                        last_err = exc
                        time.sleep(delay * (attempt + 1))
                        continue
                    raise
            return {
                "status": "busy",
                "message": f"AutoCAD Live window busy (RPC_E_CALL_REJECTED) after {max_attempts} retries: {last_err}",
            }
        return wrapper
    return decorator


def wait_for_acad_quiescent(
    acad_app: Any,
    doc: Any,
    status_file: Optional[Path] = None,
    max_wait_sec: float = 10.0,
) -> Tuple[bool, Optional[str]]:
    """Waits until AutoCAD finishes executing commands (CMDACTIVE == 0, IsQuiescent, and status_file written)."""
    time.sleep(0.35)  # Initial pause to allow SendCommand to enter the Windows message loop
    start = time.time()
    status_content = None
    while time.time() - start < max_wait_sec:
        try:
            cmd_active = doc.GetVariable("CMDACTIVE")
            state = acad_app.GetAcadState()
            is_quiescent = (cmd_active == 0 and getattr(state, "IsQuiescent", True))
            if status_file and status_file.exists():
                try:
                    status_content = status_file.read_text(encoding="utf-8-sig").strip()
                except Exception:
                    pass
            if is_quiescent and (status_file is None or status_content is not None):
                return True, status_content
        except Exception:
            pass
        time.sleep(0.2)
    return False, status_content


# ============================================================================
# V1.4.8 / V1.4.9 Extensions: Spatial BBox, Field Projection, Takeoffs, Diff, Critique & Plot
# ============================================================================
ENTITY_TYPE_NORMALIZATION: Dict[str, List[str]] = {
    "LAYERS": ["LAYER"],
    "LAYER": ["LAYER"],
    "BLOCKS": ["BLOCK", "INSERT"],
    "BLOCK": ["BLOCK", "INSERT"],
    "INSERTS": ["INSERT", "BLOCK"],
    "INSERT": ["INSERT", "BLOCK"],
    "TEXTS": ["TEXT", "MTEXT"],
    "TEXT": ["TEXT"],
    "MTEXTS": ["MTEXT"],
    "MTEXT": ["MTEXT"],
    "CURVES": ["LINE", "LWPOLYLINE", "CIRCLE", "ARC", "CURVES"],
    "CURVE": ["LINE", "LWPOLYLINE", "CIRCLE", "ARC", "CURVES"],
    "LINES": ["LINE", "CURVES"],
    "LINE": ["LINE", "CURVES"],
    "POLYLINES": ["LWPOLYLINE", "CURVES"],
    "POLYLINE": ["LWPOLYLINE", "CURVES"],
    "LWPOLYLINES": ["LWPOLYLINE", "CURVES"],
    "LWPOLYLINE": ["LWPOLYLINE", "CURVES"],
    "CIRCLES": ["CIRCLE", "CURVES"],
    "CIRCLE": ["CIRCLE", "CURVES"],
    "ARCS": ["ARC", "CURVES"],
    "ARC": ["ARC", "CURVES"],
    "TAKEOFF": ["TAKEOFF", "LAYER", "INSERT", "BLOCK", "LINE", "LWPOLYLINE", "CIRCLE", "ARC", "CURVES"],
    "METRICS": ["METRICS", "LAYER", "INSERT", "BLOCK", "LINE", "LWPOLYLINE", "CIRCLE", "ARC", "CURVES"],
}


def normalize_entity_types(raw_types: Optional[List[str]]) -> Optional[List[str]]:
    """Normalizes plural or alias entity types (e.g. LAYERS -> LAYER, BLOCKS -> BLOCK, INSERT, TEXTS -> TEXT, MTEXT, TAKEOFF -> CURVES)."""
    if not raw_types:
        return None
    normalized: List[str] = []
    for item in raw_types:
        key = item.strip().upper()
        if key in ENTITY_TYPE_NORMALIZATION:
            for mapped in ENTITY_TYPE_NORMALIZATION[key]:
                if mapped not in normalized:
                    normalized.append(mapped)
        elif key not in normalized:
            normalized.append(key)
    return normalized


UNTRUSTED_TAG_RE = re.compile(
    r"^<untrusted_cad_data[^>]*>(.*)</untrusted_cad_data>$",
    re.DOTALL,
)


def unwrap_untrusted_cad_text(val: Any) -> str:
    """Extracts the raw unescaped string from an <untrusted_cad_data> wrapper for internal diff/takeoff comparisons."""
    if val is None:
        return ""
    s = str(val)
    m = UNTRUSTED_TAG_RE.match(s.strip())
    if m:
        return html.unescape(m.group(1))
    return s


def validate_and_format_bbox(
    bbox: Optional[List[float]],
) -> Optional[Tuple[float, float, float, float]]:
    """
    Validates a spatial bounding box [xmin, ymin, xmax, ymax] in WCS coordinates.
    Returns normalized (min_x, min_y, max_x, max_y) tuple or None if bbox is None.
    """
    if bbox is None:
        return None
    if not isinstance(bbox, (list, tuple)) or len(bbox) != 4:
        raise ValueError(
            "Invalid bbox: must be a 4-element list [xmin, ymin, xmax, ymax] of numeric WCS coordinates."
        )
    coords: List[float] = []
    for v in bbox:
        if isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(float(v)):
            raise ValueError(f"Invalid numeric coordinate in bbox: {v!r}")
        coords.append(float(v))
    x1, y1, x2, y2 = coords
    return (min(x1, x2), min(y1, y2), max(x1, x2), max(y1, y2))


def filter_entities_by_bbox(
    data: Dict[str, Any],
    bbox_tuple: Optional[Tuple[float, float, float, float]],
) -> Dict[str, Any]:
    """Filters texts, blocks, and curves in data whose WCS pos [x, y, ...] lies within bbox_tuple."""
    if not bbox_tuple:
        return data
    xmin, ymin, xmax, ymax = bbox_tuple

    def _in_box(item: Dict[str, Any]) -> bool:
        pos = item.get("pos")
        if not isinstance(pos, (list, tuple)) or len(pos) < 2:
            return False
        try:
            x, y = float(pos[0]), float(pos[1])
            return (xmin <= x <= xmax) and (ymin <= y <= ymax)
        except (TypeError, ValueError):
            return False

    out = dict(data)
    out["bbox_applied"] = [xmin, ymin, xmax, ymax]
    if "texts" in out and isinstance(out["texts"], list):
        prev_len = len(out["texts"])
        out["texts"] = [t for t in out["texts"] if _in_box(t)]
        if len(out["texts"]) != prev_len or "total_texts_matched" not in out:
            out["total_texts_matched"] = len(out["texts"])
            out["texts_truncated"] = False
    if "blocks" in out and isinstance(out["blocks"], list):
        prev_len = len(out["blocks"])
        out["blocks"] = [b for b in out["blocks"] if _in_box(b)]
        if len(out["blocks"]) != prev_len or "total_blocks_matched" not in out:
            out["total_blocks_matched"] = len(out["blocks"])
            out["blocks_truncated"] = False
    if "curves" in out and isinstance(out["curves"], list):
        prev_len = len(out["curves"])
        def _curve_in_box(c: Dict[str, Any]) -> bool:
            if _in_box(c):
                return True
            end_pos = c.get("end_pos")
            if isinstance(end_pos, (list, tuple)) and len(end_pos) >= 2:
                try:
                    ex, ey = float(end_pos[0]), float(end_pos[1])
                    return (xmin <= ex <= xmax) and (ymin <= ey <= ymax)
                except (TypeError, ValueError):
                    return False
            return False
        out["curves"] = [c for c in out["curves"] if _curve_in_box(c)]
        if len(out["curves"]) != prev_len or "total_curves_matched" not in out:
            out["total_curves_matched"] = len(out["curves"])
            out["curves_truncated"] = False
    return out


def project_entity_fields(
    data: Dict[str, Any],
    fields: Optional[List[str]],
) -> Dict[str, Any]:
    """
    Projects only the requested fields onto each entity dictionary in layers, texts, blocks, and curves
    to reduce token footprint on dense drawings.
    """
    if not fields:
        return data
    allowed = {str(f).strip().lower() for f in fields if str(f).strip()}
    if not allowed:
        return data
    out = dict(data)
    for collection_key in ("layers", "texts", "blocks", "curves"):
        if collection_key in out and isinstance(out[collection_key], list):
            out[collection_key] = [
                {k: v for k, v in item.items() if k.lower() in allowed}
                for item in out[collection_key]
                if isinstance(item, dict)
            ]
    out["projected_fields"] = sorted(allowed)
    return out


def compute_takeoff_summary(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Computes metric takeoffs (linear lengths, closed areas by layer, and block BOM schedule)
    from inspected CAD data.
    """
    by_layer: Dict[str, Dict[str, Any]] = {}

    def _ensure_layer(lyr_name: str) -> Dict[str, Any]:
        if lyr_name not in by_layer:
            by_layer[lyr_name] = {
                "layer": lyr_name,
                "curve_count": 0,
                "total_length": 0.0,
                "total_closed_area": 0.0,
                "block_count": 0,
                "text_count": 0,
                "entity_types": {},
            }
        return by_layer[lyr_name]

    total_length = 0.0
    total_area = 0.0
    curves = data.get("curves", []) if isinstance(data.get("curves"), list) else []
    for c in curves:
        lyr = str(c.get("layer", "0"))
        etype = str(c.get("type", "CURVE")).upper()
        length_val = float(c.get("length", 0.0) or 0.0)
        area_val = float(c.get("area", 0.0) or 0.0) if c.get("closed") else 0.0
        entry = _ensure_layer(lyr)
        entry["curve_count"] += 1
        entry["total_length"] = round(entry["total_length"] + length_val, 4)
        entry["total_closed_area"] = round(entry["total_closed_area"] + area_val, 4)
        entry["entity_types"][etype] = entry["entity_types"].get(etype, 0) + 1
        total_length += length_val
        total_area += area_val

    blocks = data.get("blocks", []) if isinstance(data.get("blocks"), list) else []
    bom_groups: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
    for b in blocks:
        lyr = str(b.get("layer", "0"))
        eff = str(b.get("effective_name") or b.get("name") or "UNKNOWN")
        sp = str(b.get("space", "model"))
        entry = _ensure_layer(lyr)
        entry["block_count"] += 1
        entry["entity_types"]["INSERT"] = entry["entity_types"].get("INSERT", 0) + 1

        key = (eff, lyr, sp)
        if key not in bom_groups:
            sample_attrs: Dict[str, str] = {}
            for att in b.get("attributes", []) if isinstance(b.get("attributes"), list) else []:
                if isinstance(att, dict) and "tag" in att:
                    sample_attrs[str(att["tag"])] = unwrap_untrusted_cad_text(att.get("val", ""))
            bom_groups[key] = {
                "effective_name": eff,
                "layer": lyr,
                "space": sp,
                "count": 0,
                "sample_attributes": sample_attrs,
            }
        bom_groups[key]["count"] += 1

    texts = data.get("texts", []) if isinstance(data.get("texts"), list) else []
    for t in texts:
        lyr = str(t.get("layer", "0"))
        etype = str(t.get("type", "TEXT")).upper()
        entry = _ensure_layer(lyr)
        entry["text_count"] += 1
        entry["entity_types"][etype] = entry["entity_types"].get(etype, 0) + 1

    sorted_bom = sorted(
        bom_groups.values(),
        key=lambda item: (-item["count"], item["effective_name"], item["layer"]),
    )
    sorted_layers = {k: by_layer[k] for k in sorted(by_layer.keys())}

    return {
        "by_layer": sorted_layers,
        "blocks_bom": sorted_bom,
        "is_truncated": bool(data.get("curves_truncated") or data.get("blocks_truncated")),
        "totals": {
            "total_curves": len(curves),
            "total_length": round(total_length, 4),
            "total_closed_area": round(total_area, 4),
            "total_blocks": len(blocks),
            "unique_block_definitions": len(sorted_bom),
            "total_texts": len(texts),
        },
    }


def export_takeoff_report(takeoff_summary: Dict[str, Any], export_path: str) -> str:
    """
    Exports takeoff_summary to CSV (UTF-8 with BOM for direct Excel compatibility)
    or XLSX (if openpyxl is available, with automatic fallback to CSV).
    Returns the absolute path of the exported file.
    """
    out_path = Path(export_path).resolve()
    out_path.parent.mkdir(parents=True, exist_ok=True)

    by_layer = takeoff_summary.get("by_layer", {})
    blocks_bom = takeoff_summary.get("blocks_bom", [])

    if out_path.suffix.lower() == ".xlsx":
        try:
            import openpyxl
            wb = openpyxl.Workbook()
            ws_lyr = wb.active
            ws_lyr.title = "Computo_Capas"
            ws_lyr.append(["Capa", "Curvas", "Longitud_Total", "Area_Cerrada_Total", "Bloques", "Textos"])
            for lyr_name, info in by_layer.items():
                ws_lyr.append([
                    lyr_name,
                    info.get("curve_count", 0),
                    info.get("total_length", 0.0),
                    info.get("total_closed_area", 0.0),
                    info.get("block_count", 0),
                    info.get("text_count", 0),
                ])

            ws_bom = wb.create_sheet(title="Computo_Bloques_BOM")
            ws_bom.append(["Bloque_Efectivo", "Capa", "Espacio", "Cantidad", "Atributos_Muestra"])
            for item in blocks_bom:
                attrs_str = "; ".join(f"{k}={v}" for k, v in item.get("sample_attributes", {}).items())
                ws_bom.append([
                    item.get("effective_name", ""),
                    item.get("layer", ""),
                    item.get("space", ""),
                    item.get("count", 0),
                    attrs_str,
                ])
            wb.save(str(out_path))
            return str(out_path)
        except ImportError:
            out_path = out_path.with_suffix(".csv")

    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["=== COMPUTO METRICO POR CAPA ==="])
        writer.writerow(["Capa", "Curvas", "Longitud_Total", "Area_Cerrada_Total", "Bloques", "Textos"])
        for lyr_name, info in by_layer.items():
            writer.writerow([
                lyr_name,
                info.get("curve_count", 0),
                info.get("total_length", 0.0),
                info.get("total_closed_area", 0.0),
                info.get("block_count", 0),
                info.get("text_count", 0),
            ])
        writer.writerow([])
        writer.writerow(["=== PLANILLA DE BLOQUES (BOM) ==="])
        writer.writerow(["Bloque_Efectivo", "Capa", "Espacio", "Cantidad", "Atributos_Muestra"])
        for item in blocks_bom:
            attrs_str = "; ".join(f"{k}={v}" for k, v in item.get("sample_attributes", {}).items())
            writer.writerow([
                item.get("effective_name", ""),
                item.get("layer", ""),
                item.get("space", ""),
                item.get("count", 0),
                attrs_str,
            ])
    return str(out_path)


def resolve_comparison_dwg(target_dwg: Path, compare_with_dwg: str) -> Path:
    """
    Resolves a comparison DWG path. Supports:
    - 'original_backup': resolves to <stem>-Legacy-0.dwg
    - 'latest_backup': resolves to highest numbered <stem>-Legacy-<N>.dwg
    - Explicit relative or absolute path to another .dwg file
    """
    token = compare_with_dwg.strip()
    stem = target_dwg.stem
    parent = target_dwg.parent

    if token.lower() == "original_backup":
        cand = parent / f"{stem}-Legacy-0.dwg"
        if not cand.exists():
            raise FileNotFoundError(f"Original backup not found: {cand}")
        return cand.resolve()

    if token.lower() == "latest_backup":
        backups: List[Tuple[int, Path]] = []
        pattern = re.compile(rf"^{re.escape(stem)}-Legacy-(\d+)\.dwg$", re.IGNORECASE)
        if parent.exists():
            for p in parent.iterdir():
                m = pattern.match(p.name)
                if m and p.is_file():
                    backups.append((int(m.group(1)), p.resolve()))
        if not backups:
            raise FileNotFoundError(f"No -Legacy-N.dwg backups found for {target_dwg.name} in {parent}")
        backups.sort(key=lambda x: x[0], reverse=True)
        return backups[0][1]

    cand_path = Path(token)
    if not cand_path.is_absolute():
        if (parent / cand_path).exists():
            cand_path = parent / cand_path
    cand_path = cand_path.resolve()
    if not cand_path.exists():
        raise FileNotFoundError(f"Comparison DWG file not found: {cand_path}")
    return cand_path


def _point_distance(p1: Any, p2: Any) -> float:
    if not isinstance(p1, (list, tuple)) or not isinstance(p2, (list, tuple)):
        return 0.0
    if len(p1) < 2 or len(p2) < 2:
        return 0.0
    dx = float(p2[0]) - float(p1[0])
    dy = float(p2[1]) - float(p1[1])
    dz = (float(p2[2]) - float(p1[2])) if (len(p1) >= 3 and len(p2) >= 3) else 0.0
    return math.sqrt(dx * dx + dy * dy + dz * dz)


def compute_drawing_diff(
    baseline_data: Dict[str, Any],
    current_data: Dict[str, Any],
    pos_tol: float = 1e-3,
) -> Dict[str, Any]:
    """
    Computes a deterministic structural and geometric diff between two inspected CAD datasets
    (layers, texts, blocks, and optional curves).
    """
    # 1. Layers diff
    base_layers = {
        str(l["name"]): l for l in baseline_data.get("layers", []) if isinstance(l, dict) and "name" in l
    }
    curr_layers = {
        str(l["name"]): l for l in current_data.get("layers", []) if isinstance(l, dict) and "name" in l
    }
    layers_added = sorted(k for k in curr_layers if k not in base_layers)
    layers_removed = sorted(k for k in base_layers if k not in curr_layers)
    layers_modified = []
    for k in sorted(set(base_layers).intersection(curr_layers)):
        bl, cl = base_layers[k], curr_layers[k]
        if bl.get("color") != cl.get("color") or bl.get("status") != cl.get("status"):
            layers_modified.append({
                "name": k,
                "old_color": bl.get("color"),
                "new_color": cl.get("color"),
                "old_status": bl.get("status"),
                "new_status": cl.get("status"),
            })

    # 2. Texts diff (by handle)
    base_texts = {
        str(t["handle"]): t for t in baseline_data.get("texts", []) if isinstance(t, dict) and "handle" in t
    }
    curr_texts = {
        str(t["handle"]): t for t in current_data.get("texts", []) if isinstance(t, dict) and "handle" in t
    }
    texts_added = [curr_texts[h] for h in sorted(curr_texts) if h not in base_texts]
    texts_removed = [base_texts[h] for h in sorted(base_texts) if h not in curr_texts]
    texts_modified = []
    texts_moved = []
    for h in sorted(set(base_texts).intersection(curr_texts)):
        bt, ct = base_texts[h], curr_texts[h]
        old_txt = unwrap_untrusted_cad_text(bt.get("text", ""))
        new_txt = unwrap_untrusted_cad_text(ct.get("text", ""))
        old_lyr = bt.get("layer")
        new_lyr = ct.get("layer")
        if old_txt != new_txt or old_lyr != new_lyr:
            texts_modified.append({
                "handle": h,
                "old_text": bt.get("text"),
                "new_text": ct.get("text"),
                "old_layer": old_lyr,
                "new_layer": new_lyr,
            })
        dist = _point_distance(bt.get("pos"), ct.get("pos"))
        if dist > pos_tol:
            texts_moved.append({
                "handle": h,
                "text": ct.get("text"),
                "old_pos": bt.get("pos"),
                "new_pos": ct.get("pos"),
                "distance": round(dist, 4),
            })

    # 3. Blocks diff (by handle)
    base_blocks = {
        str(b["handle"]): b for b in baseline_data.get("blocks", []) if isinstance(b, dict) and "handle" in b
    }
    curr_blocks = {
        str(b["handle"]): b for b in current_data.get("blocks", []) if isinstance(b, dict) and "handle" in b
    }
    blocks_added = [curr_blocks[h] for h in sorted(curr_blocks) if h not in base_blocks]
    blocks_removed = [base_blocks[h] for h in sorted(base_blocks) if h not in curr_blocks]
    blocks_moved = []
    blocks_modified = []
    for h in sorted(set(base_blocks).intersection(curr_blocks)):
        bb, cb = base_blocks[h], curr_blocks[h]
        dist = _point_distance(bb.get("pos"), cb.get("pos"))
        eff_name = cb.get("effective_name") or cb.get("name")
        if dist > pos_tol:
            blocks_moved.append({
                "handle": h,
                "effective_name": eff_name,
                "old_pos": bb.get("pos"),
                "new_pos": cb.get("pos"),
                "distance": round(dist, 4),
            })
        old_attrs = {
            str(a.get("tag")): unwrap_untrusted_cad_text(a.get("val", ""))
            for a in (bb.get("attributes") or [])
            if isinstance(a, dict)
        }
        new_attrs = {
            str(a.get("tag")): unwrap_untrusted_cad_text(a.get("val", ""))
            for a in (cb.get("attributes") or [])
            if isinstance(a, dict)
        }
        old_eff = bb.get("effective_name") or bb.get("name")
        if old_attrs != new_attrs or bb.get("layer") != cb.get("layer") or old_eff != eff_name:
            blocks_modified.append({
                "handle": h,
                "old_effective_name": old_eff,
                "new_effective_name": eff_name,
                "old_layer": bb.get("layer"),
                "new_layer": cb.get("layer"),
                "old_attributes": old_attrs,
                "new_attributes": new_attrs,
            })

    # 4. Curves / Takeoff length, area, layer, closed & position delta by handle
    base_curves = {
        str(c["handle"]): c for c in baseline_data.get("curves", []) if isinstance(c, dict) and "handle" in c
    }
    curr_curves = {
        str(c["handle"]): c for c in current_data.get("curves", []) if isinstance(c, dict) and "handle" in c
    }
    curves_added = len([h for h in curr_curves if h not in base_curves])
    curves_removed = len([h for h in base_curves if h not in curr_curves])
    curves_modified = []
    curves_moved = []
    for h in sorted(set(base_curves).intersection(curr_curves)):
        bc, cc = base_curves[h], curr_curves[h]
        old_len = float(bc.get("length", 0.0) or 0.0)
        new_len = float(cc.get("length", 0.0) or 0.0)
        old_area = float(bc.get("area", 0.0) or 0.0)
        new_area = float(cc.get("area", 0.0) or 0.0)
        if (
            abs(new_len - old_len) > pos_tol
            or abs(new_area - old_area) > pos_tol
            or bc.get("layer") != cc.get("layer")
            or bool(bc.get("closed")) != bool(cc.get("closed"))
        ):
            curves_modified.append({
                "handle": h,
                "type": cc.get("type"),
                "old_layer": bc.get("layer"),
                "new_layer": cc.get("layer"),
                "old_length": round(old_len, 4),
                "new_length": round(new_len, 4),
                "old_area": round(old_area, 4),
                "new_area": round(new_area, 4),
            })
        dist = _point_distance(bc.get("pos"), cc.get("pos"))
        if dist > pos_tol:
            curves_moved.append({
                "handle": h,
                "type": cc.get("type"),
                "old_pos": bc.get("pos"),
                "new_pos": cc.get("pos"),
                "distance": round(dist, 4),
            })

    curves_changed_total = curves_added + curves_removed + len(curves_modified) + len(curves_moved)
    total_changes = (
        len(layers_added)
        + len(layers_removed)
        + len(layers_modified)
        + len(texts_added)
        + len(texts_removed)
        + len(texts_modified)
        + len(texts_moved)
        + len(blocks_added)
        + len(blocks_removed)
        + len(blocks_moved)
        + len(blocks_modified)
        + curves_changed_total
    )

    return {
        "layers": {
            "added": layers_added,
            "removed": layers_removed,
            "modified": layers_modified,
        },
        "texts": {
            "added": texts_added,
            "removed": texts_removed,
            "modified": texts_modified,
            "moved": texts_moved,
        },
        "blocks": {
            "added": blocks_added,
            "removed": blocks_removed,
            "moved": blocks_moved,
            "modified": blocks_modified,
        },
        "curves": {
            "added_count": curves_added,
            "removed_count": curves_removed,
            "modified": curves_modified,
            "moved": curves_moved,
        },
        "summary": {
            "is_identical": total_changes == 0,
            "total_changes": total_changes,
            "layers_changed": len(layers_added) + len(layers_removed) + len(layers_modified),
            "texts_changed": len(texts_added) + len(texts_removed) + len(texts_modified) + len(texts_moved),
            "blocks_changed": len(blocks_added) + len(blocks_removed) + len(blocks_moved) + len(blocks_modified),
            "curves_changed": curves_changed_total,
        },
    }


def critique_drawing_data(data: Dict[str, Any]) -> Dict[str, Any]:
    """
    Automated CAD hygiene and quality critique on inspected drawing data:
    - Checks INSUNITS vs MEASUREMENT consistency.
    - Detects spatial outlier entities far from the drawing core.
    - Checks excessive entity placement on Layer '0' or 'DEFPOINTS'.
    - Detects duplicate overlapping texts and duplicate overlapping block insertions.
    - Detects zero-length curves.
    Returns a 0-100 quality score and structured list of issues.
    """
    issues: List[Dict[str, Any]] = []
    penalty = 0

    insunits = data.get("insunits")
    measurement = data.get("measurement")
    if insunits == 0:
        issues.append({
            "code": "INSUNITS_UNSPECIFIED",
            "severity": "warning",
            "message": "Drawing has INSUNITS=0 (Unitless). Set INSUNITS (e.g. 4=mm, 6=m) for reliable scale and Xref insertion.",
        })
        penalty += 10
    elif insunits in (1, 2) and measurement == 1:
        issues.append({
            "code": "UNITS_MISMATCH",
            "severity": "warning",
            "message": f"INSUNITS={insunits} (Imperial inches/feet) conflicts with MEASUREMENT=1 (Metric).",
        })
        penalty += 15
    elif insunits in (4, 5, 6) and measurement == 0:
        issues.append({
            "code": "UNITS_MISMATCH",
            "severity": "warning",
            "message": f"INSUNITS={insunits} (Metric) conflicts with MEASUREMENT=0 (Imperial).",
        })
        penalty += 15

    texts = data.get("texts", []) if isinstance(data.get("texts"), list) else []
    blocks = data.get("blocks", []) if isinstance(data.get("blocks"), list) else []
    curves = data.get("curves", []) if isinstance(data.get("curves"), list) else []
    all_entities = texts + blocks + curves

    # Check Layer 0 / DEFPOINTS overuse
    if all_entities:
        lyr0_count = sum(
            1 for e in all_entities if str(e.get("layer", "")).strip().upper() in ("0", "DEFPOINTS")
        )
        ratio = lyr0_count / float(len(all_entities))
        if len(all_entities) >= 5 and ratio > 0.5:
            issues.append({
                "code": "EXCESSIVE_LAYER_ZERO",
                "severity": "warning",
                "message": f"{lyr0_count}/{len(all_entities)} ({round(ratio * 100, 1)}%) inspected entities reside on Layer '0' or 'DEFPOINTS'.",
                "count": lyr0_count,
            })
            penalty += 10

    # Check overlapping duplicate texts (scoped per space and layout)
    seen_texts: Dict[Tuple[str, str, str, str, int, int], str] = {}
    dup_texts: List[Dict[str, Any]] = []
    for t in texts:
        pos = t.get("pos")
        if isinstance(pos, (list, tuple)) and len(pos) >= 2:
            raw_t = unwrap_untrusted_cad_text(t.get("text", "")).strip()
            lyr = str(t.get("layer", "0")).upper()
            sp = str(t.get("space", "model")).lower()
            lay = str(t.get("layout", "Model")).upper()
            key = (sp, lay, lyr, raw_t, int(round(float(pos[0]) * 100)), int(round(float(pos[1]) * 100)))
            hnd = str(t.get("handle", ""))
            if raw_t and key in seen_texts:
                dup_texts.append({"handle": hnd, "duplicate_of": seen_texts[key], "layer": lyr, "layout": t.get("layout", "Model"), "pos": pos})
            elif raw_t:
                seen_texts[key] = hnd
    if dup_texts:
        issues.append({
            "code": "DUPLICATE_OVERLAPPING_TEXTS",
            "severity": "warning",
            "message": f"Detected {len(dup_texts)} duplicate overlapping text entity(ies) at identical coordinates.",
            "entities": dup_texts[:20],
        })
        penalty += min(20, 5 * len(dup_texts))

    # Check overlapping duplicate blocks (scoped per space and layout)
    seen_blocks: Dict[Tuple[str, str, str, str, int, int], str] = {}
    dup_blocks: List[Dict[str, Any]] = []
    for b in blocks:
        pos = b.get("pos")
        if isinstance(pos, (list, tuple)) and len(pos) >= 2:
            eff = str(b.get("effective_name") or b.get("name") or "").upper()
            lyr = str(b.get("layer", "0")).upper()
            sp = str(b.get("space", "model")).lower()
            lay = str(b.get("layout", "Model")).upper()
            key = (sp, lay, lyr, eff, int(round(float(pos[0]) * 100)), int(round(float(pos[1]) * 100)))
            hnd = str(b.get("handle", ""))
            if eff and key in seen_blocks:
                dup_blocks.append({"handle": hnd, "duplicate_of": seen_blocks[key], "effective_name": eff, "layer": lyr, "layout": b.get("layout", "Model"), "pos": pos})
            elif eff:
                seen_blocks[key] = hnd
    if dup_blocks:
        issues.append({
            "code": "DUPLICATE_OVERLAPPING_BLOCKS",
            "severity": "warning",
            "message": f"Detected {len(dup_blocks)} duplicate overlapping block insertion(s) at identical coordinates.",
            "entities": dup_blocks[:20],
        })
        penalty += min(20, 5 * len(dup_blocks))

    # Check zero-length curves
    zero_curves = [
        {"handle": c.get("handle"), "type": c.get("type"), "layer": c.get("layer"), "pos": c.get("pos")}
        for c in curves
        if float(c.get("length", 0.0) or 0.0) <= 1e-6
    ]
    if zero_curves:
        issues.append({
            "code": "ZERO_LENGTH_GEOMETRY",
            "severity": "error",
            "message": f"Detected {len(zero_curves)} zero-length curve entity(ies).",
            "entities": zero_curves[:20],
        })
        penalty += min(25, 5 * len(zero_curves))

    # Check spatial outliers in Model Space (prevent mixing Model and Paper Space coordinates)
    pts_with_hnd: List[Tuple[float, float, str, str]] = []
    for e in all_entities:
        if str(e.get("space", "model")).lower() != "model":
            continue
        pos = e.get("pos")
        if isinstance(pos, (list, tuple)) and len(pos) >= 2:
            try:
                pts_with_hnd.append((float(pos[0]), float(pos[1]), str(e.get("handle", "")), str(e.get("layer", ""))))
            except (TypeError, ValueError):
                pass
    if len(pts_with_hnd) >= 4:
        xs = sorted(p[0] for p in pts_with_hnd)
        ys = sorted(p[1] for p in pts_with_hnd)
        med_x = xs[len(xs) // 2]
        med_y = ys[len(ys) // 2]
        dists = sorted(math.hypot(p[0] - med_x, p[1] - med_y) for p in pts_with_hnd)
        p75_dist = max(dists[(len(dists) * 3) // 4], 100.0)
        outliers = [
            {"handle": p[2], "layer": p[3], "pos": [p[0], p[1]], "distance_from_median": round(math.hypot(p[0] - med_x, p[1] - med_y), 2)}
            for p in pts_with_hnd
            if math.hypot(p[0] - med_x, p[1] - med_y) > p75_dist * 25.0 and math.hypot(p[0] - med_x, p[1] - med_y) > 5000.0
        ]
        if outliers:
            issues.append({
                "code": "SPATIAL_OUTLIER_ENTITIES",
                "severity": "warning",
                "message": f"Detected {len(outliers)} spatial outlier entity(ies) far from the drawing median cluster.",
                "entities": outliers[:10],
            })
            penalty += 15

    score = max(0, 100 - penalty)
    return {
        "score": score,
        "issue_count": len(issues),
        "issues": issues,
    }


def build_headless_plot_script(
    output_path: Path,
    layout_name: str = "Model",
    paper_size: str = "ISO full bleed A3 (420.00 x 297.00 MM)",
    ctb_style: str = "monochrome.ctb",
    orientation: str = "Landscape",
    window_bbox: Optional[List[float]] = None,
) -> str:
    """
    Builds a non-modal international AutoCAD _.-PLOT script (.scr) for accoreconsole.exe
    to export Model or Paper Space layout to PDF or PNG for visual verification.
    Supports full _Extents or a specific WCS _Window [xmin, ymin, xmax, ymax].
    """
    out_posix = output_path.resolve().as_posix()
    ext = output_path.suffix.lower()
    if ext == ".png":
        device = "PublishToWeb PNG.pc3"
        paper = "Sun Hi-Res (1600.00 x 1280.00 Pixels)"
        units_line = ""
    else:
        device = "DWG To PDF.pc3"
        paper = paper_size
        units_line = "_Millimeters\n"

    if window_bbox is not None:
        xmin, ymin, xmax, ymax = validate_and_format_bbox(window_bbox)
        area_lines = f"_Window\n{xmin},{ymin}\n{xmax},{ymax}\n"
    else:
        area_lines = "_Extents\n"

    orient_flag = "_Landscape" if orientation.strip().lower().startswith("l") else "_Portrait"
    is_model = layout_name.strip().lower() in ("model", "model_space", "modelspace")

    if is_model:
        plot_lines = (
            "FILEDIA 0\n"
            "CMDECHO 0\n"
            "BACKGROUNDPLOT 0\n"
            "_.-PLOT\n"
            "_Yes\n"
            "Model\n"
            f"{device}\n"
            f"{paper}\n"
            f"{units_line}"
            f"{orient_flag}\n"
            "_No\n"
            f"{area_lines}"
            "_Fit\n"
            "_Center\n"
            "_Yes\n"
            f"{ctb_style}\n"
            "_Yes\n"
            "_As displayed\n"
            f'"{out_posix}"\n'
            "_No\n"
            "_Yes\n"
            "_.QUIT _Y\n"
        )
    else:
        plot_lines = (
            "FILEDIA 0\n"
            "CMDECHO 0\n"
            "BACKGROUNDPLOT 0\n"
            "_.-PLOT\n"
            "_Yes\n"
            f"{layout_name}\n"
            f"{device}\n"
            f"{paper}\n"
            f"{units_line}"
            f"{orient_flag}\n"
            "_No\n"
            f"{area_lines}"
            "_Fit\n"
            "_Center\n"
            "_Yes\n"
            f"{ctb_style}\n"
            "_Yes\n"
            "_No\n"
            "_No\n"
            "_No\n"
            f'"{out_posix}"\n'
            "_No\n"
            "_Yes\n"
            "_.QUIT _Y\n"
        )
    return plot_lines


def convert_pdf_to_png(pdf_path: Path, png_path: Path, dpi: int = 200) -> bool:
    """
    Converts the first page of a plotted vector PDF to a high-resolution PNG image
    using pypdfium2 or pdf2image for Multimodal VQA inspection.
    """
    if not pdf_path.exists() or pdf_path.stat().st_size == 0:
        return False
    png_path.parent.mkdir(parents=True, exist_ok=True)
    # 1. Try pypdfium2
    try:
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(pdf_path))
        if len(pdf) > 0:
            page = pdf[0]
            scale = max(1.0, float(dpi) / 72.0)
            bitmap = page.render(scale=scale)
            pil_image = bitmap.to_pil()
            pil_image.save(str(png_path), format="PNG")
            page.close()
            pdf.close()
            return png_path.exists() and png_path.stat().st_size > 0
    except Exception:
        pass

    # 2. Fallback to pdf2image
    try:
        from pdf2image import convert_from_path
        images = convert_from_path(str(pdf_path), dpi=dpi, first_page=1, last_page=1)
        if images:
            images[0].save(str(png_path), format="PNG")
            return png_path.exists() and png_path.stat().st_size > 0
    except Exception:
        pass

    return False


# ============================================================================
# Short-Term Improvements (1.1 Parametric Primitives & 1.2 Topological Audit)
# ============================================================================
FORBIDDEN_CAD_SYMBOL_CHARS = set('<>\\/":?*|,=`\r\n\t')


def _validate_cad_symbol_name(name: str, label: str = "layer") -> str:
    if not isinstance(name, str) or not name.strip():
        raise ValueError(f"Invalid {label} name: cannot be empty.")
    cleaned = name.strip()
    if any(ch in FORBIDDEN_CAD_SYMBOL_CHARS for ch in cleaned):
        raise ValueError(
            f"Invalid {label} name '{cleaned}': contains forbidden AutoCAD symbol table characters."
        )
    return cleaned


def _validate_point_coords(pt: Any, label: str = "point") -> Tuple[float, float, float]:
    if not isinstance(pt, (list, tuple)) or len(pt) < 2 or len(pt) > 3:
        raise ValueError(f"Invalid {label}: expected [x, y] or [x, y, z] list of numbers.")
    coords: List[float] = []
    for c in pt:
        if isinstance(c, bool) or not isinstance(c, (int, float)) or not math.isfinite(float(c)):
            raise ValueError(f"Invalid coordinate {c!r} in {label}.")
        coords.append(float(c))
    if len(coords) == 2:
        coords.append(0.0)
    return (coords[0], coords[1], coords[2])


def _validate_color_index(color: Optional[int]) -> Optional[int]:
    if color is None:
        return None
    if isinstance(color, bool) or not isinstance(color, int) or not (1 <= color <= 255):
        raise ValueError(f"Invalid ACI color index {color!r}: must be an integer between 1 and 255.")
    return color


def build_primitive_lisp_call(primitive_type: str, **kwargs: Any) -> str:
    """
    Validates parameters and synthesizes a deterministic AutoLISP call to lisp-library/primitives.lsp
    without requiring the LLM to write raw AutoLISP syntax.
    """
    ptype = primitive_type.strip().lower()
    layer = _validate_cad_symbol_name(kwargs.get("layer", "0"), "layer")
    color = _validate_color_index(kwargs.get("color"))
    color_lisp = str(color) if color is not None else "nil"
    layer_escaped = layer.replace("\\", "\\\\").replace('"', '\\"')

    if ptype == "line":
        p1 = _validate_point_coords(kwargs.get("start_pt"), "start_pt")
        p2 = _validate_point_coords(kwargs.get("end_pt"), "end_pt")
        return f'(mcp-draw-line (list {p1[0]} {p1[1]} {p1[2]}) (list {p2[0]} {p2[1]} {p2[2]}) "{layer_escaped}" {color_lisp})'

    if ptype == "circle":
        cen = _validate_point_coords(kwargs.get("center"), "center")
        rad = kwargs.get("radius")
        if isinstance(rad, bool) or not isinstance(rad, (int, float)) or float(rad) <= 0.0 or not math.isfinite(float(rad)):
            raise ValueError(f"Invalid radius {rad!r}: must be a positive number.")
        return f'(mcp-draw-circle (list {cen[0]} {cen[1]} {cen[2]}) {float(rad)} "{layer_escaped}" {color_lisp})'

    if ptype == "arc":
        cen = _validate_point_coords(kwargs.get("center"), "center")
        rad = kwargs.get("radius")
        if isinstance(rad, bool) or not isinstance(rad, (int, float)) or float(rad) <= 0.0 or not math.isfinite(float(rad)):
            raise ValueError(f"Invalid radius {rad!r}: must be a positive number.")
        sa = kwargs.get("start_angle_deg")
        ea = kwargs.get("end_angle_deg")
        if isinstance(sa, bool) or not isinstance(sa, (int, float)) or isinstance(ea, bool) or not isinstance(ea, (int, float)):
            raise ValueError("start_angle_deg and end_angle_deg must be numeric angles in degrees.")
        sa_rad = round(math.radians(float(sa)), 8)
        ea_rad = round(math.radians(float(ea)), 8)
        return f'(mcp-draw-arc (list {cen[0]} {cen[1]} {cen[2]}) {float(rad)} {sa_rad} {ea_rad} "{layer_escaped}" {color_lisp})'

    if ptype == "rectangle":
        c1 = _validate_point_coords(kwargs.get("corner1"), "corner1")
        c2 = _validate_point_coords(kwargs.get("corner2"), "corner2")
        closed_lisp = "T" if kwargs.get("closed", True) else "nil"
        return f'(mcp-draw-rectangle (list {c1[0]} {c1[1]} {c1[2]}) (list {c2[0]} {c2[1]} {c2[2]}) "{layer_escaped}" {closed_lisp} {color_lisp})'

    if ptype == "lwpolyline":
        raw_pts = kwargs.get("points")
        if not isinstance(raw_pts, (list, tuple)) or len(raw_pts) < 2:
            raise ValueError("Invalid points for lwpolyline: must provide at least 2 points.")
        pts_lisp_items = []
        for idx, pt in enumerate(raw_pts):
            p = _validate_point_coords(pt, f"points[{idx}]")
            pts_lisp_items.append(f"(list {p[0]} {p[1]})")
        pts_lisp = "(list " + " ".join(pts_lisp_items) + ")"
        closed_lisp = "T" if kwargs.get("closed", False) else "nil"
        elev = float(kwargs.get("elevation", 0.0) or 0.0)
        return f'(mcp-draw-lwpolyline {pts_lisp} {closed_lisp} "{layer_escaped}" {color_lisp} {elev})'

    if ptype == "insert_block":
        blk_name = _validate_cad_symbol_name(kwargs.get("block_name", ""), "block_name")
        blk_escaped = blk_name.replace("\\", "\\\\").replace('"', '\\"')
        ins_pt = _validate_point_coords(kwargs.get("insert_pt"), "insert_pt")
        sx = float(kwargs.get("scale_x", 1.0) or 1.0)
        sy = float(kwargs.get("scale_y", 1.0) or 1.0)
        rot_deg = float(kwargs.get("rotation_deg", 0.0) or 0.0)
        if sx == 0.0 or sy == 0.0:
            raise ValueError("Block scale factors scale_x and scale_y cannot be zero.")
        rot_rad = round(math.radians(rot_deg), 8)
        return f'(mcp-insert-block-param "{blk_escaped}" (list {ins_pt[0]} {ins_pt[1]} {ins_pt[2]}) {sx} {sy} {rot_rad} "{layer_escaped}" {color_lisp})'

    raise ValueError(f"Unsupported primitive_type: {primitive_type}")


def analyze_drawing_topology(
    data: Dict[str, Any],
    filter_layers: Optional[List[str]] = None,
    near_miss_tol: float = 5.0,
    snap_tol: float = 0.5,
) -> Dict[str, Any]:
    """
    Performs topological audit on inspected CAD curves (LINE, LWPOLYLINE):
    - Detects dangling endpoints (endpoints that do not connect to any other segment within snap_tol).
    - Detects near-miss endpoints (endpoints separated by snap_tol < d <= near_miss_tol, e.g. 1-5 mm).
    - Detects collinear/overlapping duplicate lines within snap_tol on the same layer.
    - Detects unclosed polylines (DXF 70 bit 1 unset), flagging those whose start/end points are within near_miss_tol.
    """
    allowed_layers = (
        {l.strip().upper() for l in filter_layers if l.strip()}
        if filter_layers
        else None
    )
    curves = data.get("curves", []) if isinstance(data.get("curves"), list) else []

    endpoints: List[Dict[str, Any]] = []
    unclosed_polylines: List[Dict[str, Any]] = []
    collinear_duplicates: List[Dict[str, Any]] = []
    zero_length_entities: List[Dict[str, Any]] = []
    seen_segments: List[Tuple[str, Tuple[float, float], Tuple[float, float], str]] = []

    for c in curves:
        if not isinstance(c, dict):
            continue
        lyr = str(c.get("layer", "0"))
        if allowed_layers and lyr.upper() not in allowed_layers:
            continue
        etype = str(c.get("type", "")).upper()
        hnd = str(c.get("handle", ""))
        length = float(c.get("length", 0.0) or 0.0)
        pos = c.get("pos")
        end_pos = c.get("end_pos")
        closed = bool(c.get("closed", False))

        if length <= 1e-5:
            zero_length_entities.append({"handle": hnd, "type": etype, "layer": lyr, "pos": pos})
            continue

        if etype == "LWPOLYLINE" and not closed:
            gap = _point_distance(pos, end_pos) if (pos and end_pos) else None
            is_near_closed = bool(gap is not None and gap <= near_miss_tol)
            unclosed_polylines.append({
                "handle": hnd,
                "layer": lyr,
                "start_pos": pos,
                "end_pos": end_pos,
                "closure_gap": round(gap, 4) if gap is not None else None,
                "is_near_closed": is_near_closed,
            })

        if etype in ("LINE", "LWPOLYLINE") and not closed:
            if isinstance(pos, (list, tuple)) and len(pos) >= 2:
                endpoints.append({"handle": hnd, "type": etype, "layer": lyr, "end": "start", "pt": (float(pos[0]), float(pos[1]))})
            if isinstance(end_pos, (list, tuple)) and len(end_pos) >= 2:
                endpoints.append({"handle": hnd, "type": etype, "layer": lyr, "end": "end", "pt": (float(end_pos[0]), float(end_pos[1]))})

        if etype == "LINE" and isinstance(pos, (list, tuple)) and isinstance(end_pos, (list, tuple)) and len(pos) >= 2 and len(end_pos) >= 2:
            p1 = (float(pos[0]), float(pos[1]))
            p2 = (float(end_pos[0]), float(end_pos[1]))
            is_dup_of = None
            for s_lyr, sp1, sp2, s_hnd in seen_segments:
                if s_lyr == lyr.upper():
                    d_direct = math.hypot(p1[0] - sp1[0], p1[1] - sp1[1]) + math.hypot(p2[0] - sp2[0], p2[1] - sp2[1])
                    d_rev = math.hypot(p1[0] - sp2[0], p1[1] - sp2[1]) + math.hypot(p2[0] - sp1[0], p2[1] - sp1[1])
                    if min(d_direct, d_rev) <= snap_tol * 2.0:
                        is_dup_of = s_hnd
                        break
            if is_dup_of is not None:
                collinear_duplicates.append({
                    "handle": hnd,
                    "duplicate_of": is_dup_of,
                    "layer": lyr,
                    "start_pos": [p1[0], p1[1]],
                    "end_pos": [p2[0], p2[1]],
                })
            else:
                seen_segments.append((lyr.upper(), p1, p2, hnd))

    def _dist_pt_to_seg(p: Tuple[float, float], a: Tuple[float, float], b: Tuple[float, float]) -> float:
        dx, dy = b[0] - a[0], b[1] - a[1]
        l2 = dx * dx + dy * dy
        if l2 < 1e-9:
            return math.hypot(p[0] - a[0], p[1] - a[1])
        t = max(0.0, min(1.0, ((p[0] - a[0]) * dx + (p[1] - a[1]) * dy) / l2))
        proj = (a[0] + t * dx, a[1] + t * dy)
        return math.hypot(p[0] - proj[0], p[1] - proj[1])

    dangling_endpoints: List[Dict[str, Any]] = []
    near_misses: List[Dict[str, Any]] = []
    seen_near_pairs = set()

    for i, ep1 in enumerate(endpoints):
        connected = False
        p1 = ep1["pt"]
        for j, ep2 in enumerate(endpoints):
            if i == j:
                continue
            if ep1["handle"] == ep2["handle"] and ep1["end"] == ep2["end"]:
                continue
            p2 = ep2["pt"]
            dist = math.hypot(p1[0] - p2[0], p1[1] - p2[1])
            if dist <= snap_tol:
                connected = True
            elif snap_tol < dist <= near_miss_tol and ep1["handle"] != ep2["handle"]:
                pair_key = tuple(sorted((f"{ep1['handle']}:{ep1['end']}", f"{ep2['handle']}:{ep2['end']}")))
                if pair_key not in seen_near_pairs:
                    seen_near_pairs.add(pair_key)
                    near_misses.append({
                        "handle_1": ep1["handle"],
                        "handle_2": ep2["handle"],
                        "layer_1": ep1["layer"],
                        "layer_2": ep2["layer"],
                        "pos_1": [p1[0], p1[1]],
                        "pos_2": [p2[0], p2[1]],
                        "gap_distance": round(dist, 4),
                    })
        if not connected:
            # Check T-junction connection to any other segment on same layer
            for s_lyr, sp1, sp2, s_hnd in seen_segments:
                if s_lyr == str(ep1["layer"]).upper() and s_hnd != ep1["handle"]:
                    if _dist_pt_to_seg(p1, sp1, sp2) <= snap_tol:
                        connected = True
                        break
        if not connected:
            dangling_endpoints.append({
                "handle": ep1["handle"],
                "type": ep1["type"],
                "layer": ep1["layer"],
                "endpoint": ep1["end"],
                "pos": [p1[0], p1[1]],
            })

    total_defects = (
        len(dangling_endpoints)
        + len(near_misses)
        + len(collinear_duplicates)
        + len(unclosed_polylines)
        + len(zero_length_entities)
    )
    return {
        "counts": {
            "total_defects": total_defects,
            "dangling_endpoints": len(dangling_endpoints),
            "near_misses": len(near_misses),
            "collinear_duplicates": len(collinear_duplicates),
            "unclosed_polylines": len(unclosed_polylines),
            "near_closed_polylines": sum(1 for u in unclosed_polylines if u.get("is_near_closed")),
            "zero_length_entities": len(zero_length_entities),
        },
        "dangling_endpoints": dangling_endpoints[:100],
        "near_misses": near_misses[:100],
        "collinear_duplicates": collinear_duplicates[:100],
        "unclosed_polylines": unclosed_polylines[:100],
        "zero_length_entities": zero_length_entities[:100],
    }


def execute_headless_plot(
    target_dwg: Path,
    export_plot_path: str,
    layout_name: str = "Model",
    timeout_sec: int = 60,
    window_bbox: Optional[List[float]] = None,
    paper_size: str = "ISO full bleed A3 (420.00 x 297.00 MM)",
    ctb_style: str = "monochrome.ctb",
    orientation: str = "Landscape",
) -> Dict[str, Any]:
    """Executes a headless _.-PLOT in accoreconsole.exe (/readonly) to generate a PDF/PNG visual sheet."""
    out_file = Path(export_plot_path).resolve()
    out_file.parent.mkdir(parents=True, exist_ok=True)
    if out_file.exists():
        out_file.unlink(missing_ok=True)

    run_id = uuid.uuid4().hex[:8]
    scr_file = TEMP_DIR / f"plot_{run_id}.scr"
    scr_content = build_headless_plot_script(
        out_file,
        layout_name=layout_name,
        paper_size=paper_size,
        ctb_style=ctb_style,
        orientation=orientation,
        window_bbox=window_bbox,
    )
    scr_file.write_text(scr_content, encoding="cp1252", errors="replace")

    try:
        proc = subprocess.run(
            [
                ACAD_CONSOLE_PATH,
                "/i",
                str(target_dwg),
                "/s",
                str(scr_file),
                "/l",
                "en-US",
                "/readonly",
            ],
            cwd=str(AGENTE_ROOT),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=timeout_sec,
            stdin=subprocess.DEVNULL,
        )
        stdout_text = decode_accore_stdout(proc.stdout)
        if out_file.exists() and out_file.stat().st_size > 0:
            return {
                "status": "ok",
                "plot_file": str(out_file),
                "size_bytes": out_file.stat().st_size,
                "layout": layout_name,
            }
        return {
            "status": "error",
            "message": "Plot file was not generated by accoreconsole.",
            "console_tail": stdout_text[-1000:],
        }
    except Exception as exc:
        return {"status": "error", "message": f"Headless plot error: {exc}"}
    finally:
        scr_file.unlink(missing_ok=True)


# ============================================================================
# Tool 1: Inspect CAD Headless (Primary Reader with /readonly & Delimiters)
# ============================================================================


@mcp.tool()
def inspect_cad_headless(
    dwg_path: str,
    filter_layers: Optional[List[str]] = None,
    entity_types: Optional[List[str]] = None,
    space: Optional[str] = None,
    limit: int = 200,
    timeout_sec: int = 60,
    bbox: Optional[List[float]] = None,
    fields: Optional[List[str]] = None,
    include_takeoff: bool = False,
    export_takeoff_path: Optional[str] = None,
    compare_with_dwg: Optional[str] = None,
    include_critique: bool = False,
    export_plot_path: Optional[str] = None,
    plot_layout: str = "Model",
) -> Dict[str, Any]:
    """
    PRIMARY READ TOOL. Reads layers, texts (TEXT/MTEXT), block insertions, dynamic block
    effective names, attributes, and bounds from a .dwg file using accoreconsole.exe (/readonly).
    Wraps free-form user text and attribute values in <untrusted_cad_data> delimiters
    to prevent prompt injection while keeping CAD symbol identifiers clean.
    Also reports whether the DWG is currently locked/open and by which user/machine.
    """
    cleanup_stale_temp_files()
    start_time = time.time()
    target_dwg = Path(dwg_path).resolve()
    if not target_dwg.exists():
        return {"status": "error", "message": f"File not found: {target_dwg}"}
    if not Path(ACAD_CONSOLE_PATH).exists():
        return {"status": "error", "message": f"accoreconsole.exe not found at: {ACAD_CONSOLE_PATH}"}

    blocked, block_msg, lock_owner = check_dwg_file_status(target_dwg)

    run_id = uuid.uuid4().hex[:8]
    json_out = TEMP_DIR / f"inspect_{run_id}.json"
    scr_file = TEMP_DIR / f"inspect_{run_id}.scr"

    lsp_posix = INSPECTOR_LSP.as_posix()
    scr_content = (
        "FILEDIA 0\n"
        "CMDECHO 0\n"
        "SECURELOAD 0\n"
        f'(load "{lsp_posix}")\n'
        "_.QUIT _Y\n"
    )
    scr_file.write_text(scr_content, encoding="cp1252", errors="replace")

    env = os.environ.copy()
    env["MCP_JSON_OUT"] = json_out.as_posix()
    env["MCP_LIMIT"] = str(limit)
    if filter_layers:
        cleaned_layers = [x.strip() for x in filter_layers if x.strip()]
        env["MCP_FILTER_LAYERS"] = ",".join(cleaned_layers) if cleaned_layers else "*"
    else:
        env["MCP_FILTER_LAYERS"] = "*"
    normalized_types = normalize_entity_types(entity_types)
    if normalized_types:
        env["MCP_ENTITY_TYPES"] = ",".join(normalized_types)
    else:
        env["MCP_ENTITY_TYPES"] = "*"
    if space:
        s = space.strip().upper()
        if s in ("MODEL", "MODEL_SPACE", "MODELSPACE"):
            env["MCP_SPACE"] = "MODEL"
        elif s in ("PAPER", "PAPER_SPACE", "PAPERSPACE", "LAYOUT"):
            env["MCP_SPACE"] = "PAPER"
        elif s in ("*", "ALL"):
            env["MCP_SPACE"] = "ALL"
        else:
            env["MCP_SPACE"] = s
    else:
        env["MCP_SPACE"] = "ALL"

    try:
        bbox_tuple = validate_and_format_bbox(bbox)
    except ValueError as bbox_err:
        scr_file.unlink(missing_ok=True)
        return {"status": "error", "message": str(bbox_err)}

    if bbox_tuple:
        env["MCP_BBOX"] = f"{bbox_tuple[0]},{bbox_tuple[1]},{bbox_tuple[2]},{bbox_tuple[3]}"

    want_takeoff_flag = bool(
        include_takeoff
        or export_takeoff_path
        or include_critique
        or compare_with_dwg
        or (normalized_types and ("CURVES" in normalized_types or "TAKEOFF" in normalized_types or "METRICS" in normalized_types))
    )
    if want_takeoff_flag:
        env["MCP_TAKEOFF"] = "1"

    try:
        proc = subprocess.run(
            [
                ACAD_CONSOLE_PATH,
                "/i",
                str(target_dwg),
                "/s",
                str(scr_file),
                "/l",
                "en-US",
                "/readonly",
            ],
            cwd=str(AGENTE_ROOT),
            env=env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=timeout_sec,
            stdin=subprocess.DEVNULL,
        )

        elapsed = round(time.time() - start_time, 2)
        stdout_text = decode_accore_stdout(proc.stdout)

        if "[MCP_INSPECT_ERROR]" in stdout_text:
            logging.warning(f"Inspect LISP error for {target_dwg.name}: {stdout_text}")
            return {
                "status": "error",
                "message": "AutoLISP error during inspection.",
                "is_locked": blocked,
                "lock_message": block_msg,
                "lock_owner": lock_owner,
                "console_output": stdout_text[-1500:],
                "elapsed_sec": elapsed,
            }

        if json_out.exists():
            raw_bytes = json_out.read_bytes()
            data = None
            try:
                data = json.loads(raw_bytes.decode("utf-8-sig"))
            except (json.JSONDecodeError, Exception):
                try:
                    data = json.loads(raw_bytes.decode("cp1252", errors="replace"))
                except (json.JSONDecodeError, Exception) as json_err:
                    logging.warning(f"JSON decode failed for {target_dwg.name}: {json_err}")
                    return {
                        "status": "error",
                        "message": f"Failed to parse inspection JSON output: {json_err}",
                        "is_locked": blocked,
                        "lock_message": block_msg,
                        "lock_owner": lock_owner,
                        "console_output": stdout_text[-1500:],
                        "elapsed_sec": elapsed,
                    }

            if normalized_types:
                et = set(normalized_types)
                filtered_data: Dict[str, Any] = {
                    "dwgname": data.get("dwgname"),
                    "insunits": data.get("insunits"),
                    "measurement": data.get("measurement"),
                    "tilemode": data.get("tilemode"),
                    "ctab": data.get("ctab"),
                    "extmin": data.get("extmin"),
                    "extmax": data.get("extmax"),
                }
                if "LAYER" in et and "layers" in data:
                    filtered_data["layers"] = data["layers"]
                if ("TEXT" in et or "MTEXT" in et) and "texts" in data:
                    filtered_data["texts"] = [t for t in data["texts"] if t.get("type", "").upper() in et]
                    filtered_data["total_texts_matched"] = data.get("total_texts_matched")
                    filtered_data["texts_truncated"] = data.get("texts_truncated")
                if ("INSERT" in et or "BLOCK" in et) and "blocks" in data:
                    filtered_data["blocks"] = data["blocks"]
                    filtered_data["total_blocks_matched"] = data.get("total_blocks_matched")
                    filtered_data["blocks_truncated"] = data.get("blocks_truncated")
                if ("CURVES" in et or want_takeoff_flag) and "curves" in data:
                    curve_subtypes = {"LINE", "LWPOLYLINE", "CIRCLE", "ARC"} & et
                    if curve_subtypes and "TAKEOFF" not in et and "METRICS" not in et:
                        filtered_data["curves"] = [c for c in data["curves"] if str(c.get("type", "")).upper() in curve_subtypes]
                    else:
                        filtered_data["curves"] = data["curves"]
                    filtered_data["total_curves_matched"] = data.get("total_curves_matched")
                    filtered_data["curves_truncated"] = data.get("curves_truncated")
                data = filtered_data

            # B-13: Only wrap untrusted user-created free text, preserving valid CAD symbol names
            if "texts" in data:
                for item in data["texts"]:
                    if "text" in item:
                        item["text"] = wrap_untrusted_cad_text(item["text"], "drawing_text")

            if "blocks" in data:
                for blk in data["blocks"]:
                    for attr in blk.get("attributes", []):
                        if "val" in attr:
                            attr["val"] = wrap_untrusted_cad_text(attr["val"], "block_attribute")

            if bbox_tuple:
                data = filter_entities_by_bbox(data, bbox_tuple)

            if want_takeoff_flag:
                takeoff_summary = compute_takeoff_summary(data)
                data["takeoff_summary"] = takeoff_summary
                if export_takeoff_path:
                    try:
                        data["takeoff_exported_file"] = export_takeoff_report(takeoff_summary, export_takeoff_path)
                    except Exception as exp_err:
                        data["takeoff_export_error"] = str(exp_err)

            if include_critique:
                data["critique"] = critique_drawing_data(data)

            if compare_with_dwg:
                try:
                    comp_path = resolve_comparison_dwg(target_dwg, compare_with_dwg)
                    comp_res = inspect_cad_headless(
                        dwg_path=str(comp_path),
                        filter_layers=filter_layers,
                        entity_types=entity_types,
                        space=space,
                        limit=limit,
                        timeout_sec=timeout_sec,
                        bbox=bbox,
                        include_takeoff=want_takeoff_flag,
                    )
                    if comp_res.get("status") == "ok" and isinstance(comp_res.get("data"), dict):
                        data["diff"] = compute_drawing_diff(comp_res["data"], data)
                        data["compared_against"] = str(comp_path)
                    else:
                        data["diff_error"] = comp_res.get("message", "Failed to inspect comparison DWG.")
                except Exception as diff_err:
                    data["diff_error"] = str(diff_err)

            if fields:
                data = project_entity_fields(data, fields)

            if export_plot_path:
                data["plot_export"] = execute_headless_plot(
                    target_dwg,
                    export_plot_path=export_plot_path,
                    layout_name=plot_layout,
                    timeout_sec=timeout_sec,
                    window_bbox=bbox,
                )

            logging.info(f"Inspected {target_dwg.name} in {elapsed}s (blocked={blocked})")
            return {
                "status": "ok",
                "dwg": str(target_dwg),
                "is_locked": blocked,
                "lock_message": block_msg,
                "lock_owner": lock_owner,
                "elapsed_sec": elapsed,
                "data": data,
            }

        logging.warning(f"Inspect failed for {target_dwg.name}: {stdout_text}")
        return {
            "status": "error",
            "message": "JSON output was not generated by accoreconsole.",
            "is_locked": blocked,
            "lock_message": block_msg,
            "lock_owner": lock_owner,
            "console_output": stdout_text[-1500:],
            "elapsed_sec": elapsed,
        }
    except subprocess.TimeoutExpired:
        logging.error(f"Timeout ({timeout_sec}s) inspecting {target_dwg}")
        return {"status": "error", "message": f"Process timed out after {timeout_sec}s."}
    except Exception as exc:
        logging.error(f"Error inspecting {target_dwg}: {exc}")
        return {"status": "error", "message": str(exc)}
    finally:
        scr_file.unlink(missing_ok=True)
        json_out.unlink(missing_ok=True)


# ============================================================================
# Tool 2: Mutate CAD Headless (Primary Writer with Circuit Breaker & Sequential Legacy)
# ============================================================================
@mcp.tool()
def mutate_cad_headless(
    dwg_path: str,
    lisp_code: Optional[str] = "",
    skill_script: Optional[str] = None,
    create_backup: bool = True,
    verify_audit: bool = True,
    ignore_stale_lock: bool = False,
    timeout_sec: int = 90,
    export_plot_path: Optional[str] = None,
    plot_layout: str = "Model",
    window_bbox: Optional[List[float]] = None,
) -> Dict[str, Any]:
    """
    PRIMARY WRITE TOOL. Executes AutoLISP on a closed .dwg file using accoreconsole.exe.
    - Optionally loads a pre-validated skill script from .agents/skills/ or lisp-library/ via skill_script.
    - Validates LISP code structure and safety against OS/system call circuit breakers.
    - Verifies real lock state and reports active user/machine if locked.
    - Initializes (vl-load-com), enforces DIMZIN 0, and conditionally manages UCS.
    - Cleans up hanging commands (CMDACTIVE > 0) before auditing and saving.
    - Writes .lsp with UTF-8 BOM (utf-8-sig) and .scr in cp1252.
    - Creates sequential backup (<Nombre>-Legacy-0.dwg, -Legacy-1.dwg...).
    - Traps LISP runtime errors (*error* with *mcp_err_flag*), runs _.AUDIT, saves with _.QSAVE, and verifies DBMOD == 0.
    - Automatically rolls back atomically on failure so no duplicate -Legacy files accumulate.
    """
    cleanup_stale_temp_files()
    start_time = time.time()
    target_dwg = Path(dwg_path).resolve()
    if not target_dwg.exists():
        return {"status": "error", "message": f"File not found: {target_dwg}"}
    if not Path(ACAD_CONSOLE_PATH).exists():
        return {"status": "error", "message": f"accoreconsole.exe not found at: {ACAD_CONSOLE_PATH}"}

    # 1. Structural pre-validation and circuit breaker
    try:
        skill_code = resolve_skill_script_code(skill_script) if skill_script else ""
        if skill_code:
            validate_lisp_structure(skill_code)
        if lisp_code:
            validate_lisp_structure(lisp_code)
            validate_lisp_safety(lisp_code)
        combined_lisp = "\n".join(part for part in (skill_code, lisp_code or "") if part.strip())
        if not combined_lisp.strip():
            return {"status": "error", "message": "Must provide either lisp_code or skill_script."}
        combined_lisp = normalize_skill_loads(combined_lisp)
    except ValueError as val_err:
        logging.error(f"Circuit breaker triggered: {val_err}")
        return {"status": "forbidden", "message": str(val_err)}
    except Exception as script_err:
        return {"status": "error", "message": str(script_err)}

    # 2. Check lock status and identify owner
    blocked, block_msg, lock_owner = check_dwg_file_status(target_dwg)
    if blocked:
        if (
            ignore_stale_lock
            and lock_owner
            and (lock_owner.get("is_stale") or not lock_owner.get("is_kernel_locked", True))
        ):
            logging.warning(f"Bypassing non-kernel DWL lock on {target_dwg.name} as requested by ignore_stale_lock=True")
        else:
            return {"status": "locked", "message": block_msg, "lock_owner": lock_owner}

    # 3. Create sequential Legacy backup (Legacy-0, Legacy-1...) ONLY after validations
    backup_path = None
    if create_backup:
        backup_path = create_sequential_legacy_backup(target_dwg)

    run_id = uuid.uuid4().hex[:8]
    lsp_file = TEMP_DIR / f"mutate_{run_id}.lsp"
    scr_file = TEMP_DIR / f"mutate_{run_id}.scr"

    audit_cmd = '    (if (not *mcp_err_flag*) (command "_.AUDIT" "_Y"))\n' if verify_audit else ""
    wrapped_lisp = (
        "(vl-load-com)\n"
        "(setq *mcp_err_flag* nil)\n"
        '(setq mcp_old_osmode (getvar "OSMODE"))\n'
        '(setq mcp_old_dimzin (getvar "DIMZIN"))\n'
        '(setq mcp_was_world_ucs (= (getvar "WORLDUCS") 1))\n'
        '(setq mcp_ucs_pushed nil)\n'
        "(defun *error* (msg)\n"
        "  (setq *mcp_err_flag* T)\n"
        '  (if (> (logand (getvar "CMDACTIVE") 3) 0) (while (> (logand (getvar "CMDACTIVE") 3) 0) (command)))\n'
        '  (if mcp_old_osmode (setvar "OSMODE" mcp_old_osmode))\n'
        '  (if mcp_old_dimzin (setvar "DIMZIN" mcp_old_dimzin))\n'
        '  (if mcp_ucs_pushed (progn (command "_.UCS" "_Previous") (setq mcp_ucs_pushed nil)))\n'
        '  (princ (strcat "\\n[MCP_LISP_ERROR]: " (if msg msg "Unknown error") "\\n"))\n'
        "  (princ)\n"
        ")\n"
        "(progn\n"
        '  (setvar "CMDECHO" 0)\n'
        '  (setvar "FILEDIA" 0)\n'
        '  (setvar "OSMODE" 0)\n'
        '  (setvar "DIMZIN" 0)\n'
        '  (if (not mcp_was_world_ucs) (progn (command "_.UCS" "_World") (setq mcp_ucs_pushed T)))\n'
        '  (princ "\\n[MCP_MUTATE_EXEC_START]\\n")\n'
        f"{combined_lisp}\n"
        '  (princ "\\n[MCP_MUTATE_EXEC_END]\\n")\n'
        '  (if (> (logand (getvar "CMDACTIVE") 3) 0)\n'
        "    (progn\n"
        '      (while (> (logand (getvar "CMDACTIVE") 3) 0) (command))\n'
        "      (setq *mcp_err_flag* T)\n"
        '      (princ "\\n[MCP_LISP_ERROR]: Command left active (CMDACTIVE & 3 > 0)\\n")\n'
        "    )\n"
        "  )\n"
        '  (if mcp_old_osmode (setvar "OSMODE" mcp_old_osmode))\n'
        '  (if mcp_old_dimzin (setvar "DIMZIN" mcp_old_dimzin))\n'
        '  (if mcp_ucs_pushed (progn (command "_.UCS" "_Previous") (setq mcp_ucs_pushed nil)))\n'
        f"{audit_cmd}"
        "  (if (not *mcp_err_flag*)\n"
        "    (progn\n"
        '      (command "_.QSAVE")\n'
        '      (if (= (getvar "DBMOD") 0)\n'
        '        (princ "\\n[MCP_MUTATE_SUCCESS]\\n")\n'
        '        (princ "\\n[MCP_LISP_ERROR]: QSAVE did not clear DBMOD\\n")\n'
        "      )\n"
        "    )\n"
        "  )\n"
        ")\n"
        "(princ)\n"
    )
    lsp_file.write_text(wrapped_lisp, encoding="utf-8-sig")

    scr_content = (
        "FILEDIA 0\n"
        "CMDECHO 0\n"
        "SECURELOAD 0\n"
        f'(load "{lsp_file.as_posix()}")\n'
        "_.QUIT _Y\n"
    )
    scr_file.write_text(scr_content, encoding="cp1252", errors="replace")

    try:
        proc = subprocess.run(
            [ACAD_CONSOLE_PATH, "/i", str(target_dwg), "/s", str(scr_file), "/l", "en-US"],
            cwd=str(AGENTE_ROOT),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=timeout_sec,
            stdin=subprocess.DEVNULL,
        )
        elapsed = round(time.time() - start_time, 2)
        stdout_text = decode_accore_stdout(proc.stdout)

        has_lisp_error = "[MCP_LISP_ERROR]" in stdout_text
        if "; error:" in stdout_text:
            if "[MCP_MUTATE_EXEC_START]" in stdout_text:
                s_i = stdout_text.find("[MCP_MUTATE_EXEC_START]")
                e_i = stdout_text.find("[MCP_MUTATE_EXEC_END]", s_i)
                win = stdout_text[s_i:e_i] if e_i != -1 else stdout_text[s_i:]
                if "; error:" in win:
                    has_lisp_error = True
            else:
                has_lisp_error = True
        cmd_warnings = detect_console_command_warnings(stdout_text)
        success = (
            "[MCP_MUTATE_SUCCESS]" in stdout_text
            and not has_lisp_error
            and not cmd_warnings
            and proc.returncode == 0
        )

        if not success:
            restored = False
            if backup_path and backup_path.exists():
                restored = safe_restore_backup(backup_path, target_dwg, remove_redundant_backup=True)
                if restored:
                    logging.error(f"Mutation failed on {target_dwg}. Cleaned and rolled back to pre-mutation state.")
                else:
                    logging.critical(f"Failed to restore backup {backup_path} to {target_dwg}!")
            restore_note = " Rolled back drawing to pre-mutation state." if restored else ("" if not backup_path else " Warning: Backup rollback failed.")
            return {
                "status": "error",
                "message": f"LISP execution failed or did not complete cleanly.{restore_note}",
                "command_warnings": cmd_warnings,
                "console_tail": stdout_text[-1500:],
            }

        logging.info(f"Mutated {target_dwg.name} in {elapsed}s (backup: {backup_path.name if backup_path else 'None'})")
        mut_res: Dict[str, Any] = {
            "status": "ok",
            "dwg": str(target_dwg),
            "elapsed_sec": elapsed,
            "backup_created": str(backup_path) if backup_path else None,
            "command_warnings": cmd_warnings,
            "console_tail": stdout_text[-1200:],
        }
        if export_plot_path:
            mut_res["plot_export"] = execute_headless_plot(
                target_dwg,
                export_plot_path=export_plot_path,
                layout_name=plot_layout,
                timeout_sec=timeout_sec,
                window_bbox=window_bbox,
            )
        return mut_res
    except subprocess.TimeoutExpired:
        restored = False
        if backup_path and backup_path.exists():
            restored = safe_restore_backup(backup_path, target_dwg, remove_redundant_backup=True)
        restore_note = " Drawing restored." if restored else ""
        logging.error(f"Mutation timed out for {target_dwg}")
        return {"status": "error", "message": f"Mutation timed out after {timeout_sec}s.{restore_note}"}
    except Exception as exc:
        restored = False
        if backup_path and backup_path.exists():
            restored = safe_restore_backup(backup_path, target_dwg, remove_redundant_backup=True)
        restore_note = " Drawing restored." if restored else ""
        logging.error(f"Mutation error on {target_dwg}: {exc}")
        return {"status": "error", "message": f"{exc}.{restore_note}"}
    finally:
        lsp_file.unlink(missing_ok=True)
        scr_file.unlink(missing_ok=True)


# ============================================================================
# Tool 3: Execute CAD Live (Secondary GUI Mode with Circuit Breaker)
# ============================================================================
@mcp.tool()
@com_retry(max_attempts=5, delay=0.5)
def execute_cad_live(
    lisp_code: Optional[str] = None,
    skill_script: Optional[str] = None,
    inspect_mouse_selection: bool = False,
    wait_completion: bool = True,
) -> Dict[str, Any]:
    """
    SECONDARY LIVE TOOL. Connects to the currently open AutoCAD 2021 GUI window via COM.
    Inspects user mouse selection and/or injects AutoLISP enclosed in vla-StartUndoMark / vla-EndUndoMark.
    Synchronously waits for command completion via CMDACTIVE and sentinel status files.
    Protected by AutoLISP circuit breaker, *error* handlers, and deferred temp-file cleanup.
    """
    cleanup_stale_temp_files()

    try:
        skill_code = resolve_skill_script_code(skill_script) if skill_script else ""
        if skill_code:
            validate_lisp_structure(skill_code)
        if lisp_code:
            validate_lisp_structure(lisp_code)
            validate_lisp_safety(lisp_code)
        combined_lisp = "\n".join(part for part in (skill_code, lisp_code or "") if part.strip())
        if combined_lisp:
            combined_lisp = normalize_skill_loads(combined_lisp)
    except ValueError as val_err:
        return {"status": "forbidden", "message": str(val_err)}
    except Exception as script_err:
        return {"status": "error", "message": str(script_err)}

    COM_THREAD_LOCK.acquire()
    pythoncom.CoInitialize()
    try:
        try:
            acad = _get_acad_gui_app(create_if_missing=False)
        except Exception as exc:
            if _is_com_call_rejected(exc):
                raise
            return {"status": "error", "message": f"Could not attach to AutoCAD GUI: {exc}"}

        try:
            doc = acad.ActiveDocument
            if not doc:
                return {"status": "error", "message": "No active document found in AutoCAD GUI."}
        except Exception as act_err:
            if _is_com_call_rejected(act_err):
                raise
            return {"status": "error", "message": f"No active document found in AutoCAD GUI: {act_err}"}

        result: Dict[str, Any] = {
            "status": "ok",
            "active_drawing": doc.Name,
            "full_path": doc.FullName,
        }

        if inspect_mouse_selection:
            selected_items = []
            total_selected = 0
            selection_truncated = False
            try:
                sel = doc.PickfirstSelectionSet
                total_selected = sel.Count
                selection_truncated = (total_selected > 200)
                for i in range(min(total_selected, 200)):
                    ent = sel.Item(i)
                    item_info = {
                        "handle": safe_com_getattr(ent, "Handle", ""),
                        "object_name": safe_com_getattr(ent, "ObjectName", ""),
                        "layer": safe_com_getattr(ent, "Layer", ""),
                    }
                    txt_val = safe_com_getattr(ent, "TextString")
                    if txt_val is not None:
                        item_info["text"] = wrap_untrusted_cad_text(txt_val, "selected_text")
                    eff_name = safe_com_getattr(ent, "EffectiveName")
                    if eff_name:
                        item_info["effective_name"] = eff_name
                    elif safe_com_getattr(ent, "ObjectName") == "AcDbBlockReference":
                        item_info["block_name"] = safe_com_getattr(ent, "Name", "")
                    selected_items.append(item_info)
            except Exception as sel_exc:
                if _is_com_call_rejected(sel_exc):
                    raise
                result["selection_warning"] = str(sel_exc)
            result["selected_entities"] = selected_items
            result["total_selected"] = total_selected
            result["selection_truncated"] = selection_truncated

        if combined_lisp:
            run_id = uuid.uuid4().hex[:8]
            lsp_file = TEMP_DIR / f"live_{run_id}.lsp"
            status_file = TEMP_DIR / f"live_{run_id}.status"

            old_secureload = None
            try:
                old_secureload = doc.GetVariable("SECURELOAD")
            except Exception:
                old_secureload = None

            secureload_val_str = (
                str(int(old_secureload))
                if (old_secureload is not None and isinstance(old_secureload, (int, float)))
                else '(getvar "SECURELOAD")'
            )

            status_posix = status_file.as_posix()
            wrapped_live_lisp = (
                "(vl-load-com)\n"
                "(setq mcp_live_olderr *error*)\n"
                '(setq mcp_old_cmdecho    (getvar "CMDECHO"))\n'
                '(setq mcp_old_filedia    (getvar "FILEDIA"))\n'
                '(setq mcp_old_osmode     (getvar "OSMODE"))\n'
                '(setq mcp_old_dimzin     (getvar "DIMZIN"))\n'
                f'(setq mcp_old_secureload {secureload_val_str})\n'
                "(defun *error* (msg)\n"
                '  (if (> (logand (getvar "CMDACTIVE") 3) 0) (while (> (logand (getvar "CMDACTIVE") 3) 0) (command)))\n'
                '  (if mcp_old_filedia    (setvar "FILEDIA"    mcp_old_filedia))\n'
                '  (if mcp_old_cmdecho    (setvar "CMDECHO"    mcp_old_cmdecho))\n'
                '  (if mcp_old_osmode     (setvar "OSMODE"     mcp_old_osmode))\n'
                '  (if mcp_old_dimzin     (setvar "DIMZIN"     mcp_old_dimzin))\n'
                '  (if mcp_old_secureload (setvar "SECURELOAD" mcp_old_secureload))\n'
                "  (vl-catch-all-apply '(lambda () (vla-EndUndoMark (vla-get-ActiveDocument (vlax-get-acad-object)))))\n"
                "  (setq *error* mcp_live_olderr)\n"
                f'  (setq mcp_sf (open "{status_posix}" "w"))\n'
                "  (if mcp_sf\n"
                "    (progn\n"
                '      (write-line (strcat "ERROR: " (if msg msg "Unknown error")) mcp_sf)\n'
                "      (close mcp_sf)\n"
                "    )\n"
                "  )\n"
                '  (princ (strcat "\\n[MCP_LIVE_ERROR]: " (if msg msg "Unknown error") "\\n"))\n'
                "  (princ)\n"
                ")\n"
                "(progn\n"
                "  (vl-catch-all-apply '(lambda () (vla-StartUndoMark (vla-get-ActiveDocument (vlax-get-acad-object)))))\n"
                '  (setvar "CMDECHO"    0)\n'
                '  (setvar "FILEDIA"    0)\n'
                '  (setvar "OSMODE"     0)\n'
                '  (setvar "DIMZIN"     0)\n'
                '  (setvar "SECURELOAD" 0)\n'
                f"{combined_lisp}\n"
                '  (if mcp_old_filedia    (setvar "FILEDIA"    mcp_old_filedia))\n'
                '  (if mcp_old_cmdecho    (setvar "CMDECHO"    mcp_old_cmdecho))\n'
                '  (if mcp_old_osmode     (setvar "OSMODE"     mcp_old_osmode))\n'
                '  (if mcp_old_dimzin     (setvar "DIMZIN"     mcp_old_dimzin))\n'
                '  (if mcp_old_secureload (setvar "SECURELOAD" mcp_old_secureload))\n'
                "  (vl-catch-all-apply '(lambda () (vla-EndUndoMark (vla-get-ActiveDocument (vlax-get-acad-object)))))\n"
                "  (setq *error* mcp_live_olderr)\n"
                f'  (setq mcp_sf (open "{status_posix}" "w"))\n'
                "  (if mcp_sf\n"
                "    (progn\n"
                '      (write-line "OK" mcp_sf)\n'
                "      (close mcp_sf)\n"
                "    )\n"
                "  )\n"
                "  (princ)\n"
                ")\n"
                "(princ)\n"
            )
            lsp_file.write_text(wrapped_live_lisp, encoding="utf-8-sig")

            quiescent = False
            status_text = None
            try:
                try:
                    if old_secureload is not None and old_secureload != 0:
                        doc.SetVariable("SECURELOAD", 0)
                except Exception:
                    pass

                esc_prefix = ""
                try:
                    if doc.GetVariable("CMDACTIVE") > 0:
                        esc_prefix = "\x1b\x1b"
                except Exception:
                    pass

                cmd_str = f'{esc_prefix}(load "{lsp_file.as_posix()}")\n'
                doc.SendCommand(cmd_str)

                if wait_completion:
                    quiescent, status_text = wait_for_acad_quiescent(acad, doc, status_file=status_file, max_wait_sec=10.0)
                    result["finished_quiescent"] = quiescent
                    if status_text:
                        result["live_status"] = status_text
                        if status_text.startswith("ERROR:"):
                            result["status"] = "error"
                            result["message"] = f"AutoLISP live error: {status_text}"
            finally:
                if old_secureload is not None and old_secureload != 0:
                    try:
                        doc.SetVariable("SECURELOAD", old_secureload)
                    except Exception:
                        pass
                if wait_completion and quiescent:
                    lsp_file.unlink(missing_ok=True)
                    status_file.unlink(missing_ok=True)

            result["executed_live_lisp"] = True

        return result
    finally:
        try:
            pythoncom.CoUninitialize()
        finally:
            COM_THREAD_LOCK.release()

# ============================================================================
# Tool 3b: Open Drawing in Live AutoCAD GUI
# ============================================================================
@mcp.tool()
@com_retry(max_attempts=5, delay=0.5)
def open_drawing_gui(
    dwg_path: str,
    activate_if_open: bool = True,
    create_instance_if_closed: bool = True,
) -> Dict[str, Any]:
    """
    Opens a DWG file in the visible AutoCAD 2021 GUI window so the user can inspect or edit it interactively.
    If the drawing is already open, activates its window without duplicating it.
    Thread-safe and protected by STA COM locks.
    """
    target = Path(dwg_path).resolve()
    if not target.exists():
        return {"status": "error", "message": f"Drawing file not found: {target}"}
    if not target.is_file() or target.suffix.lower() not in (".dwg", ".dxf"):
        return {"status": "error", "message": f"Target must be a .dwg or .dxf file: {target}"}

    COM_THREAD_LOCK.acquire()
    pythoncom.CoInitialize()
    try:
        try:
            acad = _get_acad_gui_app(create_if_missing=create_instance_if_closed)
        except Exception as exc:
            if _is_com_call_rejected(exc):
                raise
            return {"status": "error", "message": f"Could not connect to AutoCAD GUI: {exc}"}

        abs_target_str = str(target)
        abs_target_lower = abs_target_str.lower()

        try:
            doc_count = acad.Documents.Count
        except Exception as c_err:
            if _is_com_call_rejected(c_err):
                raise
            doc_count = 0

        for i in range(doc_count):
            try:
                d = acad.Documents.Item(i)
                if d.FullName.lower() == abs_target_lower:
                    if activate_if_open:
                        d.Activate()
                    return {
                        "status": "ok",
                        "action": "activated" if activate_if_open else "already_open",
                        "drawing_name": d.Name,
                        "full_path": d.FullName,
                        "was_already_open": True,
                    }
            except Exception as d_exc:
                if _is_com_call_rejected(d_exc):
                    raise
                continue

        try:
            doc = acad.Documents.Open(abs_target_str)
            return {
                "status": "ok",
                "action": "opened",
                "drawing_name": doc.Name,
                "full_path": doc.FullName,
                "was_already_open": False,
            }
        except Exception as o_err:
            if _is_com_call_rejected(o_err):
                raise
            return {"status": "error", "message": f"Failed to open drawing in AutoCAD GUI: {o_err}"}
    finally:
        try:
            pythoncom.CoUninitialize()
        finally:
            COM_THREAD_LOCK.release()


# ============================================================================
# Tool 3c: List Open Drawings in Live AutoCAD GUI
# ============================================================================
@mcp.tool()
@com_retry(max_attempts=5, delay=0.5)
def list_open_drawings_gui() -> Dict[str, Any]:
    """
    Lists all DWG documents currently open in the active AutoCAD 2021 GUI instance.
    Reports document name, full path, saved status, and whether it is the active window.
    """
    COM_THREAD_LOCK.acquire()
    pythoncom.CoInitialize()
    try:
        try:
            acad = _get_acad_gui_app(create_if_missing=False)
        except Exception as exc:
            if _is_com_call_rejected(exc):
                raise
            return {
                "status": "ok",
                "autocad_running": False,
                "count": 0,
                "open_drawings": [],
                "message": "AutoCAD GUI is not currently running.",
            }

        docs_list = []
        active_name = None
        try:
            active_doc = getattr(acad, "ActiveDocument", None)
            if active_doc:
                active_name = getattr(active_doc, "Name", None)
        except Exception:
            pass

        try:
            count = acad.Documents.Count
        except Exception as cnt_err:
            if _is_com_call_rejected(cnt_err):
                raise
            count = 0

        for i in range(count):
            try:
                d = acad.Documents.Item(i)
                d_name = getattr(d, "Name", "")
                d_path = getattr(d, "FullName", "")
                d_saved = getattr(d, "Saved", True)
                docs_list.append({
                    "name": d_name,
                    "full_path": d_path,
                    "saved": bool(d_saved),
                    "is_active": (d_name == active_name),
                })
            except Exception as itm_err:
                if _is_com_call_rejected(itm_err):
                    raise
                docs_list.append({"error": str(itm_err)})

        return {
            "status": "ok",
            "autocad_running": True,
            "count": len(docs_list),
            "active_drawing": active_name,
            "open_drawings": docs_list,
        }
    finally:
        try:
            pythoncom.CoUninitialize()
        finally:
            COM_THREAD_LOCK.release()


# ============================================================================
# Tool 3d: Close Drawing in Live AutoCAD GUI
# ============================================================================
@mcp.tool()
@com_retry(max_attempts=5, delay=0.5)
def close_drawing_gui(
    dwg_name: Optional[str] = None,
    save: bool = True,
) -> Dict[str, Any]:
    """
    Closes a drawing in the active AutoCAD 2021 GUI window.
    If dwg_name is omitted, closes the currently active document.
    Allows specifying whether to save changes (save=True) before closing.
    """
    COM_THREAD_LOCK.acquire()
    pythoncom.CoInitialize()
    try:
        try:
            acad = _get_acad_gui_app(create_if_missing=False)
        except Exception as exc:
            if _is_com_call_rejected(exc):
                raise
            return {"status": "error", "message": f"Could not attach to AutoCAD GUI: {exc}"}

        target_doc = None
        if not dwg_name:
            try:
                target_doc = acad.ActiveDocument
            except Exception as act_err:
                if _is_com_call_rejected(act_err):
                    raise
                return {"status": "error", "message": f"No active document found in AutoCAD GUI: {act_err}"}
        else:
            lookup = dwg_name.strip().lower()
            try:
                count = acad.Documents.Count
            except Exception as cnt_err:
                if _is_com_call_rejected(cnt_err):
                    raise
                count = 0

            for i in range(count):
                try:
                    d = acad.Documents.Item(i)
                    if d.Name.lower() == lookup or d.FullName.lower() == lookup:
                        target_doc = d
                        break
                except Exception as itm_err:
                    if _is_com_call_rejected(itm_err):
                        raise
                    continue

        if not target_doc:
            return {"status": "error", "message": f"Document '{dwg_name}' not found among open AutoCAD GUI documents."}

        closed_name = getattr(target_doc, "Name", "Unknown")
        try:
            target_doc.Close(save)
            save_msg = "saved" if save else "without saving"
            return {
                "status": "ok",
                "closed_drawing": closed_name,
                "saved": save,
                "message": f"Document '{closed_name}' closed successfully ({save_msg}).",
            }
        except Exception as cls_err:
            if _is_com_call_rejected(cls_err):
                raise
            return {"status": "error", "message": f"Failed to close document '{closed_name}': {cls_err}"}
    finally:
        try:
            pythoncom.CoUninitialize()
        finally:
            COM_THREAD_LOCK.release()


# ============================================================================
# Tool 4: Save Studio Learning (Curator Tool + Human Approval Sync)
# ============================================================================
@mcp.tool()
def save_studio_learning(
    learning_type: str,
    title: str,
    summary_for_memory: str,
    skill_slug: Optional[str] = None,
    skill_instructions_md: Optional[str] = None,
    lisp_script_content: Optional[str] = None,
    confirmed_by_user: bool = False,
    push_to_github: bool = False,
) -> Dict[str, Any]:
    """
    CURATOR TOOL. Saves validated studio criteria into MEMORY.md and/or creates a reusable
    Skill in .agents/skills/<skill_slug>/.
    Requires confirmed_by_user=True before committing and pushing to prevent memory contamination.
    """
    effective_slug = skill_slug or (title if learning_type.lower() == "skill" else None)

    if not confirmed_by_user:
        return {
            "status": "pending_confirmation",
            "message": "Changes not saved. Human-in-the-loop confirmation is required before modifying MEMORY.md or pushing to Git.",
            "proposed_learning": {
                "type": learning_type,
                "title": title,
                "summary": summary_for_memory,
                "skill_slug": effective_slug,
            },
        }

    # Validate any script content against the circuit breaker (file write strictly prohibited)
    if lisp_script_content:
        try:
            validate_lisp_structure(lisp_script_content)
            validate_lisp_safety(
                lisp_script_content,
                allow_trusted_file_write=False,
            )
        except ValueError as val_err:
            return {"status": "forbidden", "message": str(val_err)}

    # Pre-pull with rebase before reading and modifying MEMORY.md to prevent Home/Office divergence
    if (AGENTE_ROOT / ".git").exists():
        env = {
            **os.environ,
            "GIT_TERMINAL_PROMPT": "0",
            "GCM_INTERACTIVE": "never",
        }
        try:
            pull_res = subprocess.run(
                ["git", "pull", "--rebase", "--autostash"],
                cwd=str(AGENTE_ROOT),
                capture_output=True,
                text=True,
                env=env,
                timeout=20,
            )
            if pull_res.returncode != 0:
                subprocess.run(
                    ["git", "rebase", "--abort"],
                    cwd=str(AGENTE_ROOT),
                    capture_output=True,
                    text=True,
                    env=env,
                    timeout=10,
                )
                logging.warning("Pre-save git pull --rebase encountered conflicts and was aborted cleanly.")
        except Exception as exc:
            logging.warning(f"Pre-save git pull skipped: {exc}")

    memory_path = AGENTE_ROOT / "MEMORY.md"
    mem_text = memory_path.read_text(encoding="utf-8") if memory_path.exists() else "# Memoria Operativa\n"

    if len(mem_text.encode("utf-8")) > 22000:
        logging.warning("MEMORY.md exceeds 22KB budget. Consider consolidating rules.")

    created_files: List[str] = []
    global_links: List[str] = []

    if learning_type.lower() == "skill":
        raw_slug = effective_slug or title
        clean_slug = re.sub(r"[^a-z0-9-]+", "-", raw_slug.lower()).strip("-")
        if not clean_slug or len(clean_slug) < 3:
            return {"status": "error", "message": f"Invalid or too short skill_slug: '{raw_slug}'"}

        skill_dir = AGENTE_ROOT / ".agents" / "skills" / clean_slug
        scripts_dir = skill_dir / "scripts"
        scripts_dir.mkdir(parents=True, exist_ok=True)

        skill_md_path = skill_dir / "SKILL.md"
        body_md = skill_instructions_md or f"# {title}\n\n{summary_for_memory}\n"
        if not body_md.startswith("---"):
            escaped_desc = json.dumps(summary_for_memory, ensure_ascii=False)
            body_md = f"---\nname: {clean_slug}\ndescription: {escaped_desc}\n---\n\n{body_md}"
        skill_md_path.write_text(body_md, encoding="utf-8")
        created_files.append(str(skill_md_path))

        if lisp_script_content:
            lsp_path = scripts_dir / f"{clean_slug}.lsp"
            lsp_path.write_text(lisp_script_content, encoding="utf-8-sig")
            created_files.append(str(lsp_path))

        global_links = sync_skill_to_global_ides(skill_dir, clean_slug)

        bullet = f"- `{clean_slug}`: {summary_for_memory}\n"
        if f"`{clean_slug}`" not in mem_text:
            marker = "## 3. "
            if marker in mem_text:
                m_idx = mem_text.index(marker)
                nl_idx = mem_text.find("\n", m_idx)
                if nl_idx != -1:
                    insert_pos = nl_idx + 1
                    mem_text = mem_text[:insert_pos] + bullet + mem_text[insert_pos:]
                else:
                    mem_text = mem_text.rstrip() + "\n" + bullet
            else:
                mem_text = mem_text.rstrip() + "\n" + bullet
    else:
        bullet = f"- **{title}:** {summary_for_memory}\n"
        if bullet.strip() not in mem_text:
            marker = "## 3. "
            if marker in mem_text:
                idx = mem_text.index(marker)
                mem_text = mem_text[:idx].rstrip() + "\n" + bullet + "\n\n" + mem_text[idx:]
            else:
                mem_text = mem_text.rstrip() + "\n" + bullet

    memory_path.write_text(mem_text, encoding="utf-8")
    created_files.append(str(memory_path))

    git_status = "Skipped"
    if push_to_github:
        git_status = git_auto_push(f"Learn [{learning_type}]: {title}")

    logging.info(f"Saved studio learning '{title}'. Git status: {git_status}")
    return {
        "status": "ok",
        "updated_files": created_files,
        "global_skill_links": global_links,
        "git_sync": git_status,
        "last_git_pull_status": LAST_GIT_SYNC_STATUS,
    }



# ============================================================================
# Short-Term Improvement 1.1: Parametric Drawing Tools (entmake via primitives.lsp)
# ============================================================================
def _run_parametric_primitive(
    dwg_path: str,
    primitive_type: str,
    create_backup: bool = True,
    timeout_sec: int = 60,
    **kwargs: Any,
) -> Dict[str, Any]:
    try:
        lisp_call = build_primitive_lisp_call(primitive_type, **kwargs)
    except ValueError as val_err:
        return {"status": "error", "message": str(val_err)}

    res = mutate_cad_headless(
        dwg_path=dwg_path,
        lisp_code=lisp_call,
        skill_script="primitives.lsp",
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )
    if res.get("status") == "ok":
        tail = res.get("console_tail", "")
        m = re.search(r"\[MCP_PRIMITIVE_HANDLE\]:([A-Za-z0-9]+)", tail)
        if m:
            res["created_handle"] = m.group(1)
        res["primitive"] = primitive_type
        res["lisp_invoked"] = lisp_call
    return res


@mcp.tool()
def draw_primitives_batch(
    dwg_path: str,
    operations: List[Dict[str, Any]],
    create_backup: bool = True,
    timeout_sec: int = 90,
    export_plot_path: Optional[str] = None,
    plot_layout: str = "Model",
    window_bbox: Optional[List[float]] = None,
) -> Dict[str, Any]:
    """
    Executes a batch of parametric drawing primitives (line, lwpolyline, circle, rectangle, arc, insert_block)
    in a single headless accoreconsole.exe transaction with a single -Legacy-N.dwg backup and returns all created entity handles.
    """
    if not isinstance(operations, list) or not operations:
        return {"status": "error", "message": "operations must be a non-empty list of primitive dicts."}
    lisp_calls: List[str] = []
    try:
        for idx, op in enumerate(operations):
            if not isinstance(op, dict) or "type" not in op:
                raise ValueError(f"Operation at index {idx} is missing required 'type' field.")
            params = {k: v for k, v in op.items() if k != "type"}
            lisp_calls.append(build_primitive_lisp_call(str(op["type"]), **params))
    except ValueError as val_err:
        return {"status": "error", "message": f"Batch validation failed: {val_err}"}

    combined_calls = "\n".join(lisp_calls)
    res = mutate_cad_headless(
        dwg_path=dwg_path,
        lisp_code=combined_calls,
        skill_script="primitives.lsp",
        create_backup=create_backup,
        timeout_sec=timeout_sec,
        export_plot_path=export_plot_path,
        plot_layout=plot_layout,
        window_bbox=window_bbox,
    )
    if res.get("status") == "ok":
        tail = res.get("console_tail", "")
        res["created_handles"] = re.findall(r"\[MCP_PRIMITIVE_HANDLE\]:([A-Za-z0-9]+)", tail)
        res["operations_count"] = len(lisp_calls)
    return res


@mcp.tool()
def draw_line(
    dwg_path: str,
    start_pt: List[float],
    end_pt: List[float],
    layer: str = "0",
    color: Optional[int] = None,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Draws a deterministic LINE entity in WCS using entmake (via lisp-library/primitives.lsp)."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="line",
        start_pt=start_pt,
        end_pt=end_pt,
        layer=layer,
        color=color,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


@mcp.tool()
def draw_lwpolyline(
    dwg_path: str,
    points: List[List[float]],
    layer: str = "0",
    closed: bool = False,
    color: Optional[int] = None,
    elevation: float = 0.0,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Draws a deterministic LWPOLYLINE entity in WCS using entmake (via lisp-library/primitives.lsp)."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="lwpolyline",
        points=points,
        layer=layer,
        closed=closed,
        color=color,
        elevation=elevation,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


@mcp.tool()
def draw_circle(
    dwg_path: str,
    center: List[float],
    radius: float,
    layer: str = "0",
    color: Optional[int] = None,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Draws a deterministic CIRCLE entity in WCS using entmake (via lisp-library/primitives.lsp)."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="circle",
        center=center,
        radius=radius,
        layer=layer,
        color=color,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


@mcp.tool()
def draw_rectangle(
    dwg_path: str,
    corner1: List[float],
    corner2: List[float],
    layer: str = "0",
    closed: bool = True,
    color: Optional[int] = None,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Draws a deterministic rectangular LWPOLYLINE in WCS using entmake (via lisp-library/primitives.lsp)."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="rectangle",
        corner1=corner1,
        corner2=corner2,
        layer=layer,
        closed=closed,
        color=color,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


@mcp.tool()
def draw_arc(
    dwg_path: str,
    center: List[float],
    radius: float,
    start_angle_deg: float,
    end_angle_deg: float,
    layer: str = "0",
    color: Optional[int] = None,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Draws a deterministic ARC entity in WCS (angles in degrees, CCW) using entmake (via lisp-library/primitives.lsp)."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="arc",
        center=center,
        radius=radius,
        start_angle_deg=start_angle_deg,
        end_angle_deg=end_angle_deg,
        layer=layer,
        color=color,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


@mcp.tool()
def insert_block_param(
    dwg_path: str,
    block_name: str,
    insert_pt: List[float],
    layer: str = "0",
    scale_x: float = 1.0,
    scale_y: float = 1.0,
    rotation_deg: float = 0.0,
    color: Optional[int] = None,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """Inserts a block reference (INSERT) in WCS with typed scale and rotation using entmake (via lisp-library/primitives.lsp). Preserves positive scale_z when mirrored."""
    return _run_parametric_primitive(
        dwg_path=dwg_path,
        primitive_type="insert_block",
        block_name=block_name,
        insert_pt=insert_pt,
        layer=layer,
        scale_x=scale_x,
        scale_y=scale_y,
        rotation_deg=rotation_deg,
        color=color,
        create_backup=create_backup,
        timeout_sec=timeout_sec,
    )


# ============================================================================
# Short-Term Improvement 1.2: Topological Audit & Automatic Geometry Cleanup
# ============================================================================
@mcp.tool()
def audit_drawing_topology(
    dwg_path: str,
    filter_layers: Optional[List[str]] = None,
    near_miss_tol: float = 5.0,
    snap_tol: float = 0.5,
    auto_fix: bool = False,
    create_backup: bool = True,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """
    TOPOLOGICAL AUDIT & GEOMETRY CLEANUP TOOL.
    Detects dangling endpoints, near-miss vertices (e.g. 1-5mm gaps), collinear duplicate lines,
    and unclosed polylines (missing DXF 70 bit 1). When auto_fix=True, executes lisp-library/audit_topology.lsp
    to remove zero-length and duplicate collinear segments and close near-closed polylines.
    """
    insp = inspect_cad_headless(
        dwg_path=dwg_path,
        filter_layers=filter_layers,
        entity_types=["CURVES"],
        include_takeoff=True,
        limit=5000,
        timeout_sec=timeout_sec,
    )
    if insp.get("status") != "ok" or not isinstance(insp.get("data"), dict):
        return insp

    topology_report = analyze_drawing_topology(
        insp["data"],
        filter_layers=filter_layers,
        near_miss_tol=near_miss_tol,
        snap_tol=snap_tol,
    )
    result: Dict[str, Any] = {
        "status": "ok",
        "dwg": insp.get("dwg"),
        "topology": topology_report,
        "auto_fix_applied": False,
    }

    if auto_fix:
        layers_str = "*"
        if filter_layers:
            cleaned = [_validate_cad_symbol_name(l, "layer") for l in filter_layers if l.strip()]
            if cleaned:
                layers_str = ",".join(cleaned)
        layers_escaped = layers_str.replace("\\", "\\\\").replace('"', '\\"')
        lisp_call = f'(mcp-audit-and-clean-topology "{layers_escaped}" {float(snap_tol)} {float(near_miss_tol)} T)'
        mut_res = mutate_cad_headless(
            dwg_path=dwg_path,
            lisp_code=lisp_call,
            skill_script="audit_topology.lsp",
            create_backup=create_backup,
            timeout_sec=timeout_sec,
        )
        result["auto_fix_applied"] = (mut_res.get("status") == "ok")
        result["mutation_result"] = mut_res

    return result


# ============================================================================
# Short-Term Improvement 1.3: Multimodal Visual Inspection Loop (VQA Preview)
# ============================================================================
@mcp.tool()
def capture_drawing_preview(
    dwg_path: Optional[str] = None,
    output_path: Optional[str] = None,
    mode: str = "headless",
    layout_name: str = "Model",
    window_bbox: Optional[List[float]] = None,
    paper_size: str = "ISO full bleed A3 (420.00 x 297.00 MM)",
    ctb_style: str = "monochrome.ctb",
    orientation: str = "Landscape",
    dpi: int = 200,
    timeout_sec: int = 60,
) -> Dict[str, Any]:
    """
    MULTIMODAL VQA PREVIEW TOOL.
    Generates a high-resolution PNG or PDF visual preview of Model Space or a Paper Space Layout.
    - Headless mode: plots via DWG To PDF.pc3 (supporting full Extents or a specific WCS window_bbox)
      and converts the resulting vector PDF to a crisp PNG via pypdfium2 / pdf2image (with fallback to PublishToWeb PNG.pc3).
    - Live mode: sends _.ZOOM _Extents (or _Window) to the active AutoCAD GUI via COM and captures the window using Pillow/win32gui.
    """
    run_id = uuid.uuid4().hex[:8]
    out_file = Path(output_path).resolve() if output_path else (TEMP_DIR / f"preview_{run_id}.png").resolve()
    out_file.parent.mkdir(parents=True, exist_ok=True)

    if window_bbox is not None:
        try:
            validate_and_format_bbox(window_bbox)
        except ValueError as err:
            return {"status": "error", "message": str(err)}

    exec_mode = mode.strip().lower()
    if exec_mode == "live":
        pythoncom.CoInitialize()
        try:
            try:
                acad = win32com.client.GetActiveObject(ACAD_PROG_ID)
            except Exception:
                acad = win32com.client.GetActiveObject(ACAD_GENERIC_PROG_ID)
            doc = acad.ActiveDocument
            if window_bbox is not None:
                xmin, ymin, xmax, ymax = validate_and_format_bbox(window_bbox)
                doc.SendCommand(f"_.ZOOM _Window {xmin},{ymin} {xmax},{ymax}\n")
            else:
                doc.SendCommand("_.ZOOM _Extents\n")
            wait_for_acad_quiescent(acad, doc, max_wait_sec=5.0)

            try:
                import win32gui
                from PIL import ImageGrab
                hwnd = int(getattr(acad, "HWND", 0))
                rect = win32gui.GetWindowRect(hwnd) if hwnd else None
                img = ImageGrab.grab(bbox=rect) if rect else ImageGrab.grab()
                img.save(str(out_file), format="PNG")
                return {
                    "status": "ok",
                    "mode": "live",
                    "preview_file": str(out_file),
                    "format": out_file.suffix.lower().lstrip("."),
                    "size_bytes": out_file.stat().st_size,
                }
            except Exception as grab_err:
                return {"status": "error", "message": f"Live window capture failed: {grab_err}"}
        except Exception as live_err:
            return {"status": "error", "message": f"Could not attach to live AutoCAD GUI: {live_err}"}
        finally:
            pythoncom.CoUninitialize()

    # Headless mode
    if not dwg_path:
        return {"status": "error", "message": "dwg_path is required when mode='headless'."}
    target_dwg = Path(dwg_path).resolve()
    if not target_dwg.exists():
        return {"status": "error", "message": f"File not found: {target_dwg}"}
    if not Path(ACAD_CONSOLE_PATH).exists():
        return {"status": "error", "message": f"accoreconsole.exe not found at: {ACAD_CONSOLE_PATH}"}

    want_png = out_file.suffix.lower() == ".png"
    temp_pdf = (TEMP_DIR / f"preview_vec_{run_id}.pdf") if want_png else out_file
    scr_file = TEMP_DIR / f"preview_{run_id}.scr"
    scr_content = build_headless_plot_script(
        output_path=temp_pdf,
        layout_name=layout_name,
        paper_size=paper_size,
        ctb_style=ctb_style,
        orientation=orientation,
        window_bbox=window_bbox,
    )
    scr_file.write_text(scr_content, encoding="cp1252", errors="replace")

    try:
        proc = subprocess.run(
            [
                ACAD_CONSOLE_PATH,
                "/i",
                str(target_dwg),
                "/s",
                str(scr_file),
                "/l",
                "en-US",
                "/readonly",
            ],
            cwd=str(AGENTE_ROOT),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=timeout_sec,
            stdin=subprocess.DEVNULL,
        )
        stdout_text = decode_accore_stdout(proc.stdout)
        if temp_pdf.exists() and temp_pdf.stat().st_size > 0:
            if want_png:
                converted = convert_pdf_to_png(temp_pdf, out_file, dpi=dpi)
                temp_pdf.unlink(missing_ok=True)
                if converted and out_file.exists():
                    return {
                        "status": "ok",
                        "mode": "headless_pdf_to_png",
                        "preview_file": str(out_file),
                        "format": "png",
                        "dpi": dpi,
                        "layout": layout_name,
                        "size_bytes": out_file.stat().st_size,
                    }
                # Fallback to direct PublishToWeb PNG.pc3 plot if PDF renderer is not installed
                fallback_res = execute_headless_plot(
                    target_dwg,
                    export_plot_path=str(out_file),
                    layout_name=layout_name,
                    timeout_sec=timeout_sec,
                    window_bbox=window_bbox,
                    paper_size=paper_size,
                    ctb_style=ctb_style,
                    orientation=orientation,
                )
                return fallback_res
            return {
                "status": "ok",
                "mode": "headless_pdf",
                "preview_file": str(out_file),
                "format": "pdf",
                "layout": layout_name,
                "size_bytes": out_file.stat().st_size,
            }
        return {
            "status": "error",
            "message": "Preview plot file was not generated by accoreconsole.",
            "console_tail": stdout_text[-1000:],
        }
    except Exception as exc:
        return {"status": "error", "message": f"Headless preview capture failed: {exc}"}
    finally:
        scr_file.unlink(missing_ok=True)
        if want_png and temp_pdf.exists():
            temp_pdf.unlink(missing_ok=True)


if __name__ == "__main__":
    threading.Thread(target=_background_git_pull, daemon=True).start()
    mcp.run()
